from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
APP_DIR = ROOT / "arum-alpha"
OUT = ROOT / "ARUM_ALPHA_Final_Project_Report.docx"
ASSET_DIR = ROOT / "report_assets"
DEPLOYED_URL = "https://plateau-mineral-estimation.vercel.app/"
DESKTOP_SCREENSHOT = ASSET_DIR / "deployed_dashboard_desktop.png"
MOBILE_SCREENSHOT = ASSET_DIR / "deployed_dashboard_mobile.png"
STUDY_AREA_LOCATOR = ROOT / "study_area_locator.png"
UNIJOS_LOGO = ROOT / "unijos_logo.png"
TABLE_TITLES = [
    "Abbreviations used in the report",
    "Project scope and boundaries",
    "Core concepts used in the project",
    "Related methods and relevance to the project",
    "Study area and radiometric dataset summary",
    "Technology stack and component roles",
    "System architecture layers",
    "System workflow from data loading to visualization",
    "User roles and use cases",
    "Project methodology and implementation stages",
    "Feature engineering indicators",
    "Model evaluation metrics and validation methods",
    "Application programming interface endpoints",
    "Source code organization",
    "Production deployment evidence",
    "Observed project results",
    "Testing and validation summary",
    "Functional test cases",
    "Requirement traceability matrix",
    "Project setup and running commands",
    "Estimated project budget",
    "Suggested full project timeline",
]
FIGURE_TITLES = [
    "Locator view of the Plateau State project study area",
    "Desktop view of the deployed ARUM ALPHA dashboard",
    "Mobile view of the deployed ARUM ALPHA dashboard",
]
TABLE_COUNTER = 0

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(0, 0, 0)
MUTED = RGBColor(89, 89, 89)
LIGHT_FILL = "F2F4F7"
BORDER = "B8C2CC"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_TOP_BOTTOM = 80
CELL_MARGIN_START_END = 120


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_format(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_paragraph(doc: Document, text: str = "", *, style=None, bold=False, italic=False,
                  color=INK, size=11, align=None, before=0, after=6, line=1.10):
    p = doc.add_paragraph(style=style)
    set_paragraph_format(p, before=before, after=after, line=line)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_format(p, before=0, after=8, line=1.167)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run(text)
    return p


def add_numbered(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_format(p, before=0, after=8, line=1.167)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.add_run(text)
    return p


def set_cell_text(cell, text, *, bold=False, color=INK, size=10.5, align=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.text = ""
    set_paragraph_format(p, before=0, after=0, line=1.10)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold)


def set_table_geometry(table, widths_dxa: Sequence[int]):
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = borders.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)

    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in [
        ("top", CELL_MARGIN_TOP_BOTTOM),
        ("bottom", CELL_MARGIN_TOP_BOTTOM),
        ("start", CELL_MARGIN_START_END),
        ("end", CELL_MARGIN_START_END),
    ]:
        margin = margins.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]],
              widths_dxa: Sequence[int], *, font_size=9.5, numbered=True):
    global TABLE_COUNTER
    if numbered:
        TABLE_COUNTER += 1
        title = TABLE_TITLES[TABLE_COUNTER - 1] if TABLE_COUNTER <= len(TABLE_TITLES) else "Supporting project table"
        add_paragraph(
            doc,
            f"Table {TABLE_COUNTER}: {title}",
            bold=True,
            color=DARK_BLUE,
            size=9.5,
            before=4,
            after=3,
            line=1.0,
        )
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_geometry(table, widths_dxa)

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, LIGHT_FILL)
        set_cell_text(cell, header, bold=True, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if len(str(text)) <= 14 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[idx], str(text), size=font_size, align=align)

    add_paragraph(doc, "", after=4)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("ARUM ALPHA Project Report")
    set_run_font(run, size=9, color=MUTED)


def add_toc_field(paragraph):
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    run._r.append(instr)

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)

    placeholder = paragraph.add_run("Table of contents will update when opened in Microsoft Word.")
    set_run_font(placeholder, size=10, color=MUTED, italic=True)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def paragraph_bottom_border(paragraph, color="2E74B5", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def add_figure(doc: Document, image_path: Path, caption: str, *, width_inches=6.25):
    if not image_path.exists():
        add_paragraph(
            doc,
            f"Screenshot unavailable: {image_path.name}",
            italic=True,
            color=MUTED,
            size=10,
        )
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, before=6, after=3, line=1.0)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    add_paragraph(
        doc,
        caption,
        italic=True,
        color=MUTED,
        size=9.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=0,
        after=10,
        line=1.0,
    )


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    header = section.header
    header.paragraphs[0].text = ""
    hp = header.paragraphs[0]
    set_paragraph_format(hp, before=0, after=2, line=1.0)
    r = hp.add_run("ARUM ALPHA Project Report")
    set_run_font(r, size=9, color=MUTED, bold=True)
    paragraph_bottom_border(hp, color="D7DBE2", size="4")

    footer = section.footer
    footer.paragraphs[0].text = ""
    add_page_number(footer.paragraphs[0])

    props = doc.core_properties
    props.title = "ARUM ALPHA School Project Report"
    props.subject = "AI-assisted mineral estimation for Bukuru, Jos Plateau"
    props.author = "ARUM ALPHA AJIJI"
    props.comments = "Generated from the local project repository and proposal slides."


def add_cover(doc: Document):
    add_paragraph(doc, "SCHOOL PROJECT REPORT", bold=True, color=BLUE, size=12, after=18,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    if UNIJOS_LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(p, before=0, after=14, line=1.0)
        p.add_run().add_picture(str(UNIJOS_LOGO), width=Inches(0.9))
    title = (
        "DEVELOPMENT OF A MACHINE LEARNING BASED SYSTEM FOR ORE GRADE AND "
        "MINERAL DISTRIBUTION PREDICTION IN DU, BUKURU, JOS SOUTH LGA, "
        "PLATEAU STATE, NIGERIA"
    )
    add_paragraph(doc, title, bold=True, color=INK, size=20, after=8, line=1.10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "ARUM ALPHA: AI-Powered Mineral Estimation System", color=MUTED,
                  size=13, after=24, align=WD_ALIGN_PARAGRAPH.CENTER)

    rows = [
        ("Student", "ARUM ALPHA AJIJI"),
        ("Registration Number", "UJ/2018/EN/0456"),
        ("Department", "Department of Mining Engineering, University of Jos"),
        ("Supervisor", "Mr. Joshua D. Joro"),
        ("Study Area", "Du/Bukuru, Jos South LGA, Plateau State, Nigeria"),
        ("Production Link", DEPLOYED_URL),
        ("Project Type", "Final year school project documentation"),
        ("Prepared", "June 24, 2026"),
    ]
    add_table(doc, ["Field", "Details"], rows, [2300, 7060], font_size=10, numbered=False)

    p = add_paragraph(
        doc,
        "This document was prepared from the local ARUM ALPHA proposal slides, "
        "README, source code, and Sheet 168 Naraguta metadata.",
        italic=True,
        color=MUTED,
        size=10.5,
        before=10,
        after=18,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    paragraph_bottom_border(p, color="2E74B5", size="6")
    doc.add_page_break()


def add_contents(doc: Document):
    add_heading(doc, "Table of Contents", 1)
    add_paragraph(
        doc,
        "The table below is generated from the report headings and updated in Microsoft Word.",
        italic=True,
        color=MUTED,
        size=10,
        after=4,
    )
    add_toc_field(doc.add_paragraph())
    add_heading(doc, "List of Tables", 1)
    add_paragraph(
        doc,
        "The tables below are referenced within the report body so that each tabulated item is connected to the discussion.",
        italic=True,
        color=MUTED,
        size=10,
        after=4,
    )
    for idx, title in enumerate(TABLE_TITLES, start=1):
        add_paragraph(doc, f"Table {idx}: {title}", size=10, after=2, line=1.0)
    add_heading(doc, "List of Figures", 1)
    add_paragraph(
        doc,
        "The figures below are cited in the report body where each visual supports the discussion.",
        italic=True,
        color=MUTED,
        size=10,
        after=4,
    )
    for idx, title in enumerate(FIGURE_TITLES, start=1):
        add_paragraph(doc, f"Figure {idx}: {title}", size=10, after=2, line=1.0)
    doc.add_page_break()


def add_abstract(doc: Document):
    add_heading(doc, "Abstract", 1)
    add_paragraph(
        doc,
        "This project develops ARUM ALPHA, an AI-assisted web system for ore grade "
        "estimation and mineral distribution modelling in Du/Bukuru, Jos South LGA, "
        "Plateau State. The system focuses on cassiterite or tin mineralization within "
        "the Jos Plateau Younger Granite province, using radiometric potassium, thorium, "
        "and uranium grid data from Sheet 168 Naraguta as the principal local dataset.",
    )
    add_paragraph(
        doc,
        "The implemented prototype combines geospatial data processing, feature "
        "engineering, interactive mapping, AI-based geological interpretation, and "
        "dashboard visualization. It is built with Next.js, React, TypeScript, Leaflet, "
        "Recharts, Tailwind CSS, and a Groq Llama 3.3 70B integration with a rule-based "
        "fallback. The system lets a user select or search a location, retrieve nearby "
        "radiometric readings, estimate tin potential, inspect risk and confidence "
        "values, and export the result as JSON (ARUM ALPHA Project Repository, 2026).",
    )
    add_paragraph(
        doc,
        "The repository represents a working prototype toward the broader machine "
        "learning objective proposed for the project. A fully validated final model "
        "would require labelled assay or production grade data, additional geochemical "
        "and geophysical layers, and formal comparison against Random Forest, XGBoost, "
        "ordinary kriging, inverse distance weighting, and other baseline approaches "
        "(Breiman, 2001; Chen and Guestrin, 2016; Goovaerts, 1997).",
    )


def add_preliminary_sections(doc: Document):
    add_heading(doc, "Acknowledgement", 1)
    add_paragraph(
        doc,
        "I acknowledge the support of the Department of Mining Engineering, University "
        "of Jos, and the guidance of my project supervisor, Mr. Joshua D. Joro, in the "
        "development of this work. I also acknowledge the usefulness of open-source "
        "software tools, public map data, and modern web development frameworks that "
        "made it possible to build a functional prototype for mineral estimation and "
        "distribution visualization."
    )
    add_paragraph(
        doc,
        "The project benefited from the availability of radiometric grid data for Sheet "
        "168 Naraguta and from previous knowledge of tin mineralization within the Jos "
        "Plateau Younger Granite province. These resources provided the background "
        "needed to connect geoscience interpretation with a practical software system "
        "(Sheet 168 Naraguta Radiometric Metadata, 2021)."
    )

    add_heading(doc, "List of Abbreviations", 1)
    add_paragraph(
        doc,
        "Table 1 defines the abbreviations used throughout the report so that later "
        "sections can refer to technical terms consistently.",
    )
    abbreviation_rows = [
        ("AI", "Artificial Intelligence"),
        ("API", "Application Programming Interface"),
        ("DEM", "Digital Elevation Model"),
        ("GRD", "Geosoft Grid"),
        ("K", "Potassium"),
        ("ML", "Machine Learning"),
        ("NGSA", "Nigerian Geological Survey Agency"),
        ("RMSE", "Root Mean Square Error"),
        ("SnO2", "Tin oxide / cassiterite grade expression"),
        ("Th", "Thorium"),
        ("U", "Uranium"),
        ("UTM", "Universal Transverse Mercator"),
        ("WGS84", "World Geodetic System 1984"),
    ]
    add_table(doc, ["Abbreviation", "Meaning"], abbreviation_rows, [1800, 7560], font_size=9.5)


def add_intro_problem_objectives(doc: Document):
    add_heading(doc, "1. Introduction", 1)
    add_paragraph(
        doc,
        "Bukuru and neighbouring communities in Jos South occur within the Younger "
        "Granite province of the Jos Plateau, an area historically associated with "
        "cassiterite mineralization and tin mining. Traditional exploration workflows "
        "in such areas can be expensive because geological structures, hydrothermal "
        "alteration zones, and radiometric anomalies are spatially complex.",
    )
    add_paragraph(
        doc,
        "ARUM ALPHA addresses this challenge by presenting geological and radiometric "
        "information in an interactive web dashboard. The system is designed to help "
        "students, mining engineers, and exploration teams reason about likely ore "
        "grade potential before committing to detailed field verification.",
    )
    add_heading(doc, "1.1 Background of the Study", 2)
    add_paragraph(
        doc,
        "Nigeria has a long history of solid mineral exploitation, but many exploration "
        "decisions are still constrained by incomplete datasets, manual interpretation, "
        "and limited access to localized digital tools. In Plateau State, tin mining "
        "has played an important historical role, especially around Jos, Bukuru, and "
        "nearby communities. Cassiterite mineralization in the region is commonly "
        "associated with Younger Granite bodies, alteration zones, structures, and "
        "weathered or alluvial concentrations. These geological controls are not always "
        "easy to interpret from field observation alone, particularly where data are "
        "sparse or spatially irregular."
    )
    add_paragraph(
        doc,
        "Radiometric surveys are useful in such settings because potassium, thorium, "
        "and uranium responses can indicate lithological variation, alteration, and "
        "geochemical contrast. When these datasets are combined with coordinates and "
        "visualized on a map, they provide a stronger basis for target generation. "
        "However, raw geophysical data are difficult for non-specialists to use unless "
        "they are processed and presented in a clear decision-support interface "
        "(Sheet 168 Naraguta Radiometric Metadata, 2021)."
    )
    add_paragraph(
        doc,
        "Machine learning and AI-assisted interpretation provide an opportunity to "
        "move beyond static maps. Instead of only displaying data, a system can compute "
        "features such as K/(Th+U), compare nearby readings, estimate relative mineral "
        "potential, and return recommendations for field verification. ARUM ALPHA was "
        "therefore developed as a bridge between mining engineering knowledge and "
        "modern software engineering practice (Rodriguez-Galiano et al., 2015)."
    )

    add_heading(doc, "1.2 Project Motivation", 2)
    add_paragraph(
        doc,
        "The motivation for this project is the need for a localized, low-cost, and "
        "accessible exploration support tool for tin mineralization in Bukuru and the "
        "Jos Plateau area. A student, supervisor, or exploration officer should be able "
        "to open a browser, inspect the study area, select a location, and receive a "
        "structured interpretation of that location's mineral potential. This does not "
        "replace detailed fieldwork; it makes fieldwork more focused (ARUM ALPHA "
        "Proposal Slides, 2026; ARUM ALPHA Project Proposal, 2026)."
    )
    add_paragraph(
        doc,
        "The deployed prototype also demonstrates that mining engineering projects can "
        "produce usable digital tools, not only written analysis. By combining a public "
        "web interface with geoscience reasoning, the project shows how Nigerian solid "
        "minerals research can adopt modern data systems while remaining grounded in "
        "local geological conditions (ARUM ALPHA Project Repository, 2026; ARUM ALPHA "
        "Deployed Application, 2026)."
    )

    add_heading(doc, "1.3 Source of the Project Idea and Creation", 2)
    add_paragraph(
        doc,
        "The idea for ARUM ALPHA originated from the original project proposal, which "
        "identified the need for a machine-learning-based system for ore grade and "
        "mineral distribution prediction in Du/Bukuru, Jos South LGA, Plateau State "
        "(ARUM ALPHA Proposal Slides, 2026; ARUM ALPHA Project Proposal, 2026). The "
        "proposal framed the project as a mining engineering problem: how to reduce "
        "uncertainty in early-stage tin prospectivity assessment by combining local "
        "geoscience data, predictive modelling, and clear visual communication."
    )
    add_paragraph(
        doc,
        "The creation of the implemented system was guided by the available ARUM ALPHA "
        "README, source code, dashboard components, API routes, and deployment files. "
        "These project artifacts provided the practical software basis for converting "
        "the proposal idea into a deployed web dashboard with coordinate input, map "
        "visualization, K/Th/U feature handling, AI-assisted interpretation, and JSON "
        "export (ARUM ALPHA Project Repository, 2026)."
    )
    add_paragraph(
        doc,
        "The technical data foundation for the project came from the Sheet 168 Naraguta "
        "radiometric metadata and grid files for potassium, thorium, and uranium. "
        "Those files gave the project a real local dataset and helped connect the "
        "software idea to the Bukuru/Naraguta geological setting instead of leaving it "
        "as a generic mapping application (Sheet 168 Naraguta Radiometric Metadata, "
        "2021)."
    )

    add_heading(doc, "2. Problem Statement", 1)
    add_paragraph(
        doc,
        "The central problem is the absence of a localized and interactive predictive "
        "system that connects radiometric data, geological reasoning, and mineral "
        "potential estimation for the Bukuru/Jos Plateau tin field. Existing data may "
        "be available in technical formats, but without processing and visualization "
        "they do not directly support quick decision-making."
    )
    for item in [
        "Ore grade estimation in Bukuru is uncertain because the local geology is complex and mineralization may be controlled by multiple non-linear factors.",
        "Traditional interpolation methods such as ordinary kriging and inverse distance weighting can struggle when assumptions of stationarity or simple spatial continuity do not hold (Goovaerts, 1997).",
        "Fragmented geochemical, geophysical, and spatial datasets make it difficult to form a single operational picture for exploration decisions.",
        "There is a need for a localized, data-driven dashboard tailored to tin mineralization in the Jos Plateau Younger Granite setting.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. Aim and Objectives", 1)
    add_paragraph(
        doc,
        "The aim is to design and implement a machine learning based predictive system "
        "for ore grade estimation and mineral distribution modelling in Bukuru, Jos, "
        "Nigeria."
    )
    for item in [
        "Compile and preprocess multi-source datasets relevant to cassiterite exploration.",
        "Engineer useful geochemical and radiometric features from potassium, thorium, uranium, coordinates, and surrounding context.",
        "Develop predictive models that can estimate grade, confidence, and exploration risk for selected locations.",
        "Compare machine learning outputs with traditional geostatistical approaches when labelled validation data become available.",
        "Design a visualization dashboard that supports map-based exploration decisions and technical recommendations.",
    ]:
        add_numbered(doc, item)

    add_heading(doc, "3.1 Research Questions", 2)
    for item in [
        "How can radiometric K, Th, and U grid data be transformed into useful features for tin potential interpretation?",
        "How can a web dashboard improve the communication of mineral distribution patterns in the Bukuru study area?",
        "How can AI-assisted analysis be combined with rule-based geoscience logic to provide explainable exploration guidance?",
        "What additional data would be required before the prototype can become a fully validated ore grade prediction model?",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3.2 Research Assumptions", 2)
    for item in [
        "Elevated thorium and favourable Th/U relationships can indicate geological conditions associated with mineralized granite systems.",
        "Radiometric response alone is not sufficient for reserve estimation but can support early-stage target ranking.",
        "A dashboard-based workflow improves communication of spatial information when compared with isolated spreadsheet or static map outputs.",
        "Model confidence must be treated as a decision-support value until field assay data are available for supervised validation.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4. Project Scope and Requirements", 1)
    add_paragraph(
        doc,
        "The project scope covers the design, implementation, deployment, and "
        "documentation of a prototype mineral estimation platform for the Bukuru/Jos "
        "Plateau study area. The work is not presented as a certified reserve "
        "estimation system; rather, it is a decision-support prototype that combines "
        "radiometric data, geological reasoning, and dashboard visualization. Table 2 "
        "summarizes the project scope and boundary conditions."
    )
    scope_rows = [
        ("In scope", "Web dashboard, map visualization, coordinate-based estimation, K/Th/U feature extraction, AI/rule-based interpretation, production deployment, and report documentation."),
        ("Partly implemented", "Supervised machine learning workflow is represented conceptually; final training requires labelled assay or production grade data."),
        ("Out of scope", "Certified mineral reserve reporting, drilling design approval, laboratory assay generation, and legal mining lease evaluation."),
    ]
    add_table(doc, ["Scope Area", "Description"], scope_rows, [2200, 7160])

    add_heading(doc, "Functional Requirements", 2)
    for item in [
        "The system shall display the Sheet 168 Naraguta study area on an interactive map.",
        "The system shall allow users to enter latitude and longitude values or select a location on the map.",
        "The system shall retrieve radiometric context for the selected location when it lies inside the study bounds.",
        "The system shall return predicted tin potential, confidence, risk level, interpretation, and recommendations.",
        "The system shall allow exported estimation results in JSON format for later review.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Non-Functional Requirements", 2)
    for item in [
        "The interface should be readable on desktop and mobile screen sizes.",
        "The dashboard should load quickly enough for classroom demonstration and project defence use.",
        "The API should return clear errors when coordinates or estimation requests are invalid.",
        "The project should be deployable to Vercel with environment variables separated from source code.",
        "The report should clearly separate prototype results from scientifically validated mineral reserve estimates.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5. Significance of the Study", 1)
    add_paragraph(
        doc,
        "The project contributes a localized digital approach to mineral prospectivity "
        "mapping for Bukuru and the Jos Plateau. It supports exploration efficiency by "
        "reducing blind sampling, providing early-stage target ranking, and presenting "
        "geoscience information in a form that is easier to interpret."
    )
    for item in [
        "Mining engineers can use the dashboard to compare potential targets and prioritize field verification.",
        "Exploration companies can reduce early-stage uncertainty before committing drilling or sampling budgets.",
        "Policymakers and researchers can use the prototype as a reference for sustainable, data-driven solid minerals development.",
        "The project demonstrates how modern web technologies and AI reasoning can be applied to Nigerian mining problems.",
    ]:
        add_bullet(doc, item)


def add_lit_study_data(doc: Document):
    add_heading(doc, "6. Literature Review and Research Gap", 1)
    add_heading(doc, "6.1 Conceptual Review", 2)
    add_paragraph(
        doc,
        "Geostatistical methods such as kriging remain important in mineral resource "
        "estimation because they model spatial continuity and can produce interpolated "
        "grade surfaces. However, their reliability depends on assumptions about the "
        "spatial structure of the data. In structurally complex terrains, machine "
        "learning models can capture nonlinear relationships among geochemical, "
        "geophysical, topographic, and spatial variables (Goovaerts, 1997; "
        "Rodriguez-Galiano et al., 2015)."
    )
    add_paragraph(
        doc,
        "The project proposal identified Random Forest and XGBoost as appropriate "
        "supervised learning candidates because they handle nonlinear feature "
        "interactions, missing values, and feature importance analysis better than "
        "simple linear regression baselines. The present repository implements the "
        "dashboard, radiometric ingestion, geochemical indices, AI interpretation, and "
        "estimation workflow needed to support that later supervised modelling phase "
        "(Breiman, 2001; Chen and Guestrin, 2016; ARUM ALPHA Project Repository, 2026). "
        "Table 3 explains the major concepts used in the project."
    )
    concept_rows = [
        ("Ore grade estimation", "The process of estimating the concentration or quality of useful mineral material at a location or within a deposit."),
        ("Mineral prospectivity", "The ranking of areas according to the likelihood that they host mineralization of economic interest."),
        ("Radiometric survey", "A geophysical method that measures natural gamma radiation associated mainly with potassium, thorium, and uranium."),
        ("Cassiterite", "Tin oxide mineral, SnO2, and the main ore mineral for tin in many granite-related deposits."),
        ("Feature engineering", "The transformation of raw data into meaningful variables such as ratios, anomaly indicators, and spatial context features."),
    ]
    add_table(doc, ["Concept", "Meaning in this Project"], concept_rows, [2200, 7160])

    add_heading(doc, "6.2 Theoretical Framework", 2)
    add_paragraph(
        doc,
        "The theoretical foundation of the project rests on three connected ideas: "
        "geological control of mineralization, radiometric expression of rocks and "
        "alteration, and predictive modelling. In the Jos Plateau, tin mineralization "
        "is commonly discussed in relation to Younger Granite intrusions and associated "
        "hydrothermal or residual concentration processes. If these processes affect "
        "the distribution of radiogenic elements, then K, Th, and U data can provide "
        "indirect evidence for mineral prospectivity."
    )
    add_paragraph(
        doc,
        "From a modelling perspective, the system assumes that mineral potential can "
        "be represented as a function of spatial position, radiometric values, derived "
        "ratios, and surrounding context. Classical geostatistics emphasizes spatial "
        "continuity and interpolation, while machine learning emphasizes pattern "
        "recognition across multiple variables. ARUM ALPHA combines these ideas in a "
        "prototype workflow by using spatial lookup and feature calculation before "
        "AI-assisted interpretation."
    )

    add_heading(doc, "6.3 Review of Related Methods", 2)
    method_rows = [
        ("Ordinary kriging", "Strong for spatial interpolation when variogram structure is reliable; may struggle where stationarity assumptions are weak (Goovaerts, 1997)."),
        ("Inverse distance weighting", "Simple and intuitive interpolation baseline; does not learn complex geological controls."),
        ("Linear regression", "Useful baseline for checking simple relationships; limited for nonlinear mineralization patterns."),
        ("Random Forest", "Handles nonlinear variables and feature importance well; proposed for later supervised model development (Breiman, 2001)."),
        ("XGBoost", "Powerful gradient boosting model suitable for tabular prediction and missing-value handling; proposed for later validation (Chen and Guestrin, 2016)."),
        ("AI-assisted interpretation", "Useful for explaining geophysical patterns in natural language; must be checked against field evidence."),
    ]
    add_paragraph(
        doc,
        "Table 4 compares the related estimation and interpretation methods discussed "
        "in the literature review.",
    )
    add_table(doc, ["Method", "Relevance to the Project"], method_rows, [2500, 6860])

    add_heading(doc, "6.4 Empirical Review and Research Gap", 2)
    for item in [
        "Few student or local prototype systems focus specifically on tin mineralization in Bukuru and the Jos Plateau.",
        "Many mineral prospectivity examples are either generic or rely on single-source datasets rather than integrated geochemical, geophysical, and spatial evidence.",
        "The region needs a practical dashboard that turns technical radiometric data into location-specific exploration guidance.",
    ]:
        add_bullet(doc, item)
    add_paragraph(
        doc,
        "The major gap addressed by this project is therefore not only prediction, but "
        "also integration. A useful system for a local mining engineering project must "
        "combine data ingestion, coordinate handling, interpretation logic, visual "
        "communication, and deployment. ARUM ALPHA contributes to this gap by presenting "
        "a working dashboard that can be demonstrated, inspected, and extended "
        "(ARUM ALPHA Project Repository, 2026)."
    )

    add_heading(doc, "7. Study Area and Dataset", 1)
    add_paragraph(
        doc,
        "The study area covers Sheet 168 Naraguta around Bukuru, Du, Naraguta, Jos, "
        "Rayfield, Gyel, Zawan, and nearby Jos South/Jos North communities. The "
        "software bounds the local radiometric study area between about 8.4999 and "
        "8.9994 degrees east longitude, and 9.5001 and 9.9991 degrees north latitude "
        "(Sheet 168 Naraguta Radiometric Metadata, 2021). Figure 1 provides the "
        "locator view used to connect the written study-area description with the "
        "mapped Plateau State project location."
    )
    add_figure(
        doc,
        STUDY_AREA_LOCATOR,
        "Figure 1: Locator view showing the Plateau State project study area used for ARUM ALPHA.",
        width_inches=5.75,
    )
    dataset_rows = [
        ("Primary data layers", "Potassium, thorium, and uranium radiometric grids"),
        ("File format", "Geosoft Grid with XML metadata"),
        ("Grid sheet", "Sheet168_Naraguta"),
        ("Projection", "WGS 84 / UTM Zone 32N, EPSG:32632"),
        ("UTM extent", "445187.5E to 499937.5E; 1050187.5N to 1105312.5N"),
        ("Geographic extent", "8.499873E to 8.999431E; 9.500149N to 9.999096N"),
        ("Dataset date in metadata", "April 16, 2021"),
        ("Application sample limit", "Up to 1,000 points returned for map visualization"),
    ]
    add_paragraph(
        doc,
        "Table 5 summarizes the source dataset, projection, extents, and map sampling "
        "details taken from the Sheet 168 Naraguta metadata.",
    )
    add_table(doc, ["Dataset Item", "Description"], dataset_rows, [2500, 6860])
    add_paragraph(
        doc,
        "The metadata confirms that the local grids were produced from Nigerian "
        "radiometric source grids using a Sheet 168 Naraguta mask. The current "
        "application parser contains a fallback for older Geosoft GRD formats, so "
        "raw grid value validation should be completed before the system is used for "
        "formal resource estimation (Sheet 168 Naraguta Radiometric Metadata, 2021)."
    )
    add_heading(doc, "7.1 Geological Setting", 2)
    add_paragraph(
        doc,
        "The Jos Plateau Younger Granite province is widely known for tin-bearing "
        "granites and associated minerals such as columbite, tantalite, wolframite, "
        "and monazite. Bukuru and neighbouring areas have experienced historic mining "
        "activity, and the landscape contains evidence of previous exploration and "
        "extraction. The geological setting is important because radiometric anomalies "
        "must be interpreted within the context of granite bodies, structures, "
        "weathering, and possible hydrothermal alteration."
    )
    add_paragraph(
        doc,
        "In this project, the study area is treated as a geospatial problem: every "
        "radiometric value has a coordinate, and every coordinate can be related to "
        "the map, local communities, and surrounding readings. This spatial framing "
        "makes it possible to move from isolated data points to a distribution model."
    )

    add_heading(doc, "7.2 Dataset Preparation Considerations", 2)
    for item in [
        "Grid metadata must be read carefully so that projected coordinates are interpreted correctly.",
        "No-data or dummy values must be excluded from statistical and modelling calculations.",
        "Coordinates must be converted from UTM to latitude/longitude before they can be displayed accurately on a web map.",
        "Values from K, Th, and U grids must be combined by matching grid positions so that every sampled point has a complete feature record.",
        "The final model should document whether each point is measured, interpolated, synthetic, or derived from an external source.",
    ]:
        add_bullet(doc, item)


def add_design_method_impl(doc: Document):
    add_heading(doc, "8. System Design and Technology Stack", 1)
    add_paragraph(
        doc,
        "ARUM ALPHA is implemented as a Next.js web application with API routes for "
        "data retrieval and mineral estimation. The frontend presents a map, search "
        "tools, coordinate input, estimation results, radiometric values, risk "
        "assessment, model metrics, and JSON export. Table 6 lists the main technology "
        "components used in the implementation (ARUM ALPHA Project Repository, 2026)."
    )
    tech_rows = [
        ("Frontend framework", "Next.js 16, React 19, TypeScript"),
        ("Styling", "Tailwind CSS 4"),
        ("Mapping", "Leaflet and React-Leaflet"),
        ("AI service", "Groq SDK using llama-3.3-70b-versatile"),
        ("Charts and icons", "Recharts and Lucide React"),
        ("Core data code", "GRD parser, UTM conversion, nearest point lookup, radius sampling"),
        ("Deployment target", "Vercel configuration included"),
    ]
    add_table(doc, ["Component", "Technology or Role"], tech_rows, [2600, 6760])

    add_heading(doc, "8.1 System Architecture", 2)
    add_paragraph(
        doc,
        "The system follows a web application architecture made up of a client-facing "
        "dashboard, server API routes, data processing utilities, and an external AI "
        "service. The client dashboard handles interaction and visualization, while "
        "the server routes load radiometric data and run estimation logic. This "
        "separation keeps the user interface responsive and allows sensitive API keys "
        "to remain outside the browser. Table 7 summarizes these architecture layers."
    )
    architecture_rows = [
        ("Presentation layer", "Dashboard, map view, prediction panel, coordinate form, search interface, legend, and result export."),
        ("Application/API layer", "Next.js route handlers for /api/data and /api/estimate."),
        ("Data layer", "Local Geosoft GRD files and XML metadata for Sheet 168 Naraguta."),
        ("Processing layer", "GRD parsing, UTM conversion, nearest-neighbour lookup, radius filtering, and ratio calculation."),
        ("AI interpretation layer", "Groq Llama model prompt plus rule-based fallback for mineral potential estimation."),
        ("Deployment layer", "Vercel hosting with production URL for public access."),
    ]
    add_table(doc, ["Layer", "Description"], architecture_rows, [2500, 6860])

    add_heading(doc, "System Workflow", 2)
    workflow_rows = [
        ("1", "Data load", "Load K, Th, and U grids or fallback synthetic data"),
        ("2", "Coordinate selection", "User searches, enters coordinates, or clicks the map"),
        ("3", "Spatial lookup", "Find nearest point and optional surrounding points within radius"),
        ("4", "Feature engineering", "Calculate K/(Th+U), Th/U, and contextual averages"),
        ("5", "Estimation", "Generate grade, confidence, risk, analysis, and recommendations"),
        ("6", "Visualization", "Render markers, heatmap overlays, result cards, and export file"),
    ]
    add_paragraph(
        doc,
        "Table 8 presents the workflow from data loading to map visualization and "
        "prediction output.",
    )
    add_table(doc, ["Step", "Process", "Output"], workflow_rows, [900, 2400, 6060], font_size=9)

    add_heading(doc, "8.2 User Roles and Use Cases", 2)
    use_case_rows = [
        ("Student/researcher", "Demonstrate the project, inspect radiometric distribution, and explain the estimation workflow."),
        ("Supervisor/reviewer", "Open the deployed dashboard, test coordinates, and assess whether the implementation matches the report."),
        ("Mining engineer", "Use the dashboard as an early-stage guide for selecting field verification targets."),
        ("Policy or planning user", "Understand how local geoscience data can be converted into a decision-support platform."),
    ]
    add_paragraph(
        doc,
        "Table 9 identifies the main users of the dashboard and the use case each "
        "user group is expected to perform.",
    )
    add_table(doc, ["User", "Use Case"], use_case_rows, [2400, 6960])

    add_heading(doc, "8.3 Data Flow Description", 2)
    add_paragraph(
        doc,
        "When the application starts, the data API loads radiometric points and returns "
        "a sampled subset for map display. When the user submits coordinates, the "
        "estimation API checks whether the location falls inside the study area. For "
        "inside-area locations, it searches for the nearest data point and optionally "
        "nearby points within the selected radius. The estimator then calculates "
        "ratios, builds a geological prompt, returns predicted grade and risk, and "
        "sends the response back to the dashboard."
    )
    for item in [
        "Input: user-selected latitude and longitude.",
        "Processing: bounds check, nearest point lookup, surrounding point selection, feature calculation, and estimation.",
        "Output: prediction object containing grade, confidence, risk, K/Th/U values, analysis, recommendations, and model metrics.",
        "Storage/export: the current prototype exports JSON from the browser rather than storing results in a database.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "9. Methodology", 1)
    add_paragraph(
        doc,
        "The methodology follows the proposal workflow: data acquisition, preprocessing, "
        "feature engineering, model development, validation, spatial prediction, and "
        "dashboard deployment. The implemented prototype covers the ingestion, feature, "
        "AI interpretation, visualization, and deployment parts of this workflow. "
        "Table 10 connects each methodology stage to the implemented project activity."
    )
    methodology = [
        ("Data acquisition", "Radiometric K, Th, and U grid files and XML metadata are stored in the repository data folder."),
        ("Preprocessing", "The parser reads grid metadata, applies WGS84 UTM Zone 32N coordinate assumptions, and converts UTM coordinates to latitude/longitude for mapping."),
        ("Feature engineering", "The system calculates K/(Th+U) and Th/U ratios, and optionally averages surrounding points within a user-selected search radius."),
        ("Prediction logic", "The main estimator prompts an AI model with geophysical context and falls back to a rule-based estimator if the API fails."),
        ("Risk classification", "Predicted grade thresholds are used to communicate low, medium, and high exploration risk in the dashboard."),
        ("Visualization", "Leaflet renders the study area, data markers, prediction markers, base layers, and heatmap overlays."),
    ]
    add_table(doc, ["Stage", "Project Implementation"], methodology, [2300, 7060])

    add_heading(doc, "9.1 Detailed Research Procedure", 2)
    for item in [
        "Review the geological and mining context of Bukuru, Du, Naraguta, and the wider Jos Plateau tin field.",
        "Collect local radiometric datasets and inspect their XML metadata for projection, extent, date, and source details.",
        "Build a parser and fallback data generation pathway so the dashboard can run during demonstrations even when raw grid parsing fails.",
        "Convert projected grid coordinates into geographic coordinates suitable for Leaflet web mapping.",
        "Design user interface controls for search, coordinate entry, map interaction, estimation display, and result export.",
        "Implement API routes that separate data loading from estimation logic.",
        "Deploy the system to Vercel and verify that the production interface loads from a public URL.",
        "Document limitations and future validation needs so the prototype is not confused with a certified reserve estimate.",
    ]:
        add_numbered(doc, item)

    add_heading(doc, "9.2 Feature Engineering", 2)
    add_paragraph(
        doc,
        "Feature engineering is a key part of the methodology because raw radiometric "
        "values are more useful when transformed into indicators that have geological "
        "meaning. The current system uses potassium, thorium, uranium, K/(Th+U), "
        "Th/U, coordinate position, and optional surrounding-point averages. These "
        "features are simple enough to explain during project defence but still "
        "capture meaningful relationships for tin prospectivity screening. Table 11 "
        "shows the interpretation use of each feature."
    )
    feature_rows = [
        ("Potassium (%)", "Helps indicate feldspar-rich or altered zones in granitic terrain."),
        ("Thorium (ppm)", "Elevated values may indicate radiometric anomaly and possible granite-related alteration."),
        ("Uranium (ppm)", "Used with thorium and potassium to describe radiometric character and ratios."),
        ("K/(Th+U)", "Balances potassium response against thorium and uranium concentration."),
        ("Th/U", "Highlights relative thorium enrichment and possible magmatic differentiation indicators."),
        ("Neighbourhood average", "Provides local context around a selected point rather than relying on one reading only."),
    ]
    add_table(doc, ["Feature", "Interpretation Use"], feature_rows, [2500, 6860])

    add_heading(doc, "9.3 Model Evaluation Plan", 2)
    add_paragraph(
        doc,
        "The deployed prototype includes demonstration metrics, but a complete "
        "scientific model requires labelled target values. When assay or production "
        "grade data become available, the proposed evaluation plan is to split the "
        "dataset into training, validation, and test sets, compare models using the "
        "same features, and report error metrics consistently. Table 12 lists the "
        "metrics proposed for future model validation."
    )
    eval_rows = [
        ("RMSE", "Measures typical prediction error with stronger penalty for large errors."),
        ("MAE", "Measures average absolute prediction error in a more interpretable way."),
        ("R2", "Indicates the proportion of grade variance explained by the model."),
        ("MAPE", "Shows average percentage error where target values are suitable for percentage comparison."),
        ("Cross-validation", "Reduces dependence on one train-test split and gives a more stable accuracy estimate."),
    ]
    add_table(doc, ["Metric/Method", "Purpose"], eval_rows, [2500, 6860])

    add_heading(doc, "10. Implementation", 1)
    add_paragraph(
        doc,
        "The main dashboard is implemented in components/Dashboard.tsx. It manages "
        "coordinate input, search results, selected point state, prediction state, "
        "loading states, map-layer toggles, and JSON result export. The map is loaded "
        "dynamically to avoid server-side rendering issues with Leaflet. Table 13 "
        "summarizes the API endpoints that support the dashboard "
        "(ARUM ALPHA Project Repository, 2026)."
    )
    endpoint_rows = [
        ("GET /api/data", "Loads radiometric data, bounds, sample points, and estimated grid dimensions for visualization."),
        ("POST /api/estimate", "Accepts latitude/longitude, finds local radiometric context, and returns prediction, analysis, recommendations, and metrics."),
        ("GET /api/estimate", "Simple health check that reports API status and loaded data point count."),
    ]
    add_table(doc, ["Endpoint", "Purpose"], endpoint_rows, [2600, 6760])
    add_heading(doc, "10.1 Source Code Organization", 2)
    code_rows = [
        ("app/page.tsx", "Loads initial data on the server and renders the dashboard."),
        ("components/Dashboard.tsx", "Coordinates user input, search, prediction state, map layers, stats, and export."),
        ("components/MapView.tsx", "Renders Leaflet map, data markers, prediction markers, base layers, and study boundary."),
        ("components/PredictionPanel.tsx", "Displays predicted grade, risk, radiometric readings, analysis, and recommendations."),
        ("lib/grdParser.ts", "Parses grid data, converts UTM coordinates, generates fallback data, and performs spatial lookup."),
        ("lib/groq.ts", "Builds AI prompts, calls the Groq model, and provides a rule-based fallback estimator."),
        ("lib/onlineData.ts", "Provides broader geological context for locations outside the Naraguta study area."),
        ("app/api/data/route.ts", "Returns sampled radiometric data and study bounds for visualization."),
        ("app/api/estimate/route.ts", "Processes estimation requests and returns prediction responses."),
    ]
    add_paragraph(
        doc,
        "Table 14 maps the most important source files to their roles in the final "
        "system implementation.",
    )
    add_table(doc, ["File", "Role in the System"], code_rows, [2700, 6660], font_size=9)

    add_heading(doc, "10.2 Estimation Logic", 2)
    add_paragraph(
        doc,
        "The estimator follows a hybrid logic. First, it uses deterministic geospatial "
        "processing to find relevant radiometric data. Second, it calculates ratios "
        "and surrounding context. Third, it sends a structured prompt to the AI model "
        "requesting a JSON response with predicted grade, confidence, risk level, "
        "analysis, and recommendations. If the AI request fails, the fallback estimator "
        "uses rule-based geoscience thresholds such as thorium greater than 20 ppm, "
        "Th/U greater than 4, and potassium greater than 3 percent."
    )
    add_paragraph(
        doc,
        "This hybrid approach was selected because it allows the dashboard to remain "
        "usable even when the AI service is unavailable. It also makes the prototype "
        "easier to explain: the AI layer provides natural-language reasoning, while "
        "the fallback layer shows the geological assumptions explicitly."
    )
    add_paragraph(
        doc,
        "A production build was executed successfully on June 23, 2026. The build "
        "compiled the app, checked TypeScript, generated static pages, and preserved "
        "dynamic server routes for /api/data and /api/estimate."
    )


def add_deployment_and_screenshots(doc: Document):
    add_heading(doc, "11. Production Deployment and Screenshots", 1)
    add_paragraph(
        doc,
        "The project has been deployed publicly on Vercel for demonstration and "
        "assessment (ARUM ALPHA Deployed Application, 2026). The production link is:"
    )
    add_paragraph(doc, DEPLOYED_URL, bold=True, color=BLUE, size=11, after=10)
    add_paragraph(
        doc,
        "The deployed system demonstrates that the application can be accessed outside "
        "the development machine, loads the study-area dashboard, and presents the "
        "Naraguta radiometric map interface through a public web address. This is "
        "important for project defence because supervisors and reviewers can inspect "
        "the implementation without running the source code locally. Table 15 records "
        "the deployment evidence used in the report, while Figures 2 and 3 provide "
        "visual evidence from the live deployed dashboard on desktop and mobile "
        "viewports."
    )
    deployment_rows = [
        ("Hosting platform", "Vercel"),
        ("Production URL", DEPLOYED_URL),
        ("Observed dashboard data count", "1,000 sampled visualization points"),
        ("Visible study area", "Bukuru, Jos Plateau"),
        ("Primary interface elements", "Location search, coordinate input, map layers, legend, and map-based radiometric point display"),
    ]
    add_table(doc, ["Deployment Item", "Evidence"], deployment_rows, [2600, 6760])

    add_heading(doc, "Desktop Dashboard Screenshot", 2)
    add_paragraph(
        doc,
        "Figure 2 shows the desktop version of the deployed dashboard. It is included "
        "to demonstrate that the production site loads the study-area map, radiometric "
        "points, coordinate controls, and dashboard interface in a wide-screen layout.",
    )
    add_figure(
        doc,
        DESKTOP_SCREENSHOT,
        "Figure 2: Desktop view of the deployed ARUM ALPHA dashboard showing the study-area map and radiometric data points.",
        width_inches=6.25,
    )

    add_heading(doc, "Mobile Dashboard Screenshot", 2)
    add_paragraph(
        doc,
        "Figure 3 confirms that the dashboard layout remains usable on a narrow "
        "viewport. The interface stacks the search, coordinate input, and map sections "
        "vertically while preserving the study-area context."
    )
    add_figure(
        doc,
        MOBILE_SCREENSHOT,
        "Figure 3: Mobile view of the deployed dashboard showing responsive layout and the mapped Naraguta data points.",
        width_inches=2.65,
    )

    add_heading(doc, "Deployment Discussion", 2)
    add_paragraph(
        doc,
        "The deployed dashboard strengthens the project by showing that the prototype "
        "is not limited to local execution. It provides a visible product outcome: a "
        "web-based mineral estimation interface that can be reviewed, tested, and "
        "extended. In a final-year project setting, this supports both the technical "
        "implementation requirement and the presentation requirement because the "
        "application can be demonstrated live."
    )


def add_results_limits_conclusion(doc: Document):
    add_heading(doc, "12. Results and Discussion", 1)
    add_paragraph(
        doc,
        "The resulting system is a working web dashboard that supports interactive "
        "mineral potential estimation. A user can search for Plateau locations such "
        "as Bukuru, Naraguta, Rayfield, Gyel, Du, Zawan, and Sabon Barki, or can click "
        "any coordinate on the map. For locations inside the Sheet 168 study area, the "
        "system uses local radiometric data. For locations outside the study area, it "
        "uses a broader geological context service and synthetic radiometric estimates."
    )
    add_paragraph(
        doc,
        "The prediction panel communicates the estimated tin grade, confidence, risk "
        "level, potassium/thorium/uranium readings, K/(Th+U) ratio, AI analysis, field "
        "recommendations, and model performance metrics. The visualization gives "
        "students and technical reviewers a fast way to inspect how radiometric "
        "anomalies relate to predicted mineral potential. Table 16 summarizes the "
        "observed outcomes from the implemented prototype."
    )
    result_rows = [
        ("Dashboard", "Implemented with controls, map, prediction card, legend, stats bar, and export."),
        ("Geospatial display", "Study boundary, base map layers, sampled radiometric points, prediction markers, and heatmap layer."),
        ("Estimation output", "Predicted grade, confidence, risk level, analysis, recommendations, and model metrics."),
        ("Build verification", "npm.cmd run build completed successfully with Next.js 16.2.4."),
    ]
    add_table(doc, ["Result Area", "Observed Outcome"], result_rows, [2400, 6960])
    add_heading(doc, "12.1 Interpretation of Project Outputs", 2)
    add_paragraph(
        doc,
        "The most important output of the project is the conversion of technical "
        "radiometric information into a usable dashboard. Instead of requiring the "
        "user to inspect raw grid files, the system presents mapped points, a study "
        "boundary, coordinate tools, and an estimation panel. This makes the project "
        "suitable for classroom demonstration, supervisor review, and early exploration "
        "discussion."
    )
    add_paragraph(
        doc,
        "Figures 2 and 3 show that the deployed system successfully renders the study "
        "area and a dense set of radiometric data points. The desktop view provides a "
        "wide map and side controls, while the mobile view confirms that the interface "
        "can stack controls and map content on a smaller screen. This is important "
        "because field users may not always work from a desktop computer."
    )
    add_heading(doc, "12.2 Educational and Technical Value", 2)
    for item in [
        "It demonstrates how a mining engineering problem can be transformed into a digital decision-support system.",
        "It shows the relationship between geophysical data, geospatial coordinates, and mineral prospectivity interpretation.",
        "It gives the student a deployable software artifact that can be tested from a production link.",
        "It creates a foundation for future supervised machine learning once labelled grade data are acquired.",
        "It encourages clearer communication between geology, mining engineering, and software development.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "13. Testing and Validation", 1)
    add_paragraph(
        doc,
        "Testing focused on confirming that the application builds, the deployed "
        "interface loads, the core routes are defined, and the generated report remains "
        "openable in Microsoft Word. Scientific validation of mineral grade prediction "
        "requires labelled field or assay data and is therefore listed as future work. "
        "Table 17 summarizes the validation areas checked during the project."
    )
    testing_rows = [
        ("Production build", "npm.cmd run build completed successfully with Next.js 16.2.4."),
        ("Deployment inspection", "The Vercel dashboard page loaded and screenshots were captured from the live URL."),
        ("Responsive interface", "Desktop and mobile screenshots confirm the dashboard is usable across different viewport widths."),
        ("Document validation", "The generated DOCX was opened through Microsoft Word automation and reported as a valid document."),
        ("Scientific validation status", "Pending labelled assay or production-grade data for supervised ML training and test-set evaluation."),
    ]
    add_table(doc, ["Validation Area", "Result"], testing_rows, [2600, 6760])
    add_heading(doc, "13.1 Test Cases", 2)
    test_case_rows = [
        ("Load dashboard", "Open production URL", "Dashboard header, controls, and map should render.", "Passed"),
        ("Display study data", "Wait for map and sampled points", "Map should show study boundary and radiometric markers.", "Passed"),
        ("Coordinate input", "Enter 9.750000, 8.750000", "System should accept numeric latitude and longitude.", "Implemented"),
        ("API data route", "Call GET /api/data", "Route should return bounds and sampled radiometric points.", "Implemented"),
        ("API estimate route", "Submit POST /api/estimate", "Route should return prediction, metrics, analysis, and recommendations.", "Implemented"),
        ("Responsive layout", "Capture mobile-width screenshot", "Interface should remain usable on narrow screens.", "Passed"),
        ("DOCX validity", "Open report with Microsoft Word", "Document should open and TOC should update.", "Passed"),
    ]
    add_paragraph(
        doc,
        "Table 18 lists the functional test cases used to confirm that the deployed "
        "prototype and generated report behave as expected.",
    )
    add_table(doc, ["Test", "Action", "Expected Result", "Status"], test_case_rows, [1700, 2300, 4160, 1200], font_size=8.8)

    add_heading(doc, "13.2 Validation Limitations", 2)
    add_paragraph(
        doc,
        "The software validation confirms that the prototype runs and presents the "
        "expected interface, but it does not prove geological accuracy. Geological "
        "validation requires comparison between model predictions and independent "
        "field evidence such as assay results, trench samples, drill holes, mine "
        "production records, or expert-validated prospectivity maps. This distinction "
        "is important because a working dashboard is a software success, while an "
        "accurate ore grade model is a geoscience validation task."
    )

    add_heading(doc, "13.3 Ethical, Safety, and Environmental Considerations", 2)
    for item in [
        "The dashboard should not be used to justify mining activity without proper permits, community consultation, and environmental assessment.",
        "Field verification should follow safety procedures for abandoned mines, pits, unstable ground, and radiation-aware sampling practices.",
        "Predictions should be presented as decision-support outputs, not guaranteed ore grades.",
        "Data sources and uncertainty should be documented to avoid misleading investors, communities, or government agencies.",
        "Sustainable exploration should minimize unnecessary land disturbance by using digital target ranking before field sampling.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "14. Limitations", 1)
    for item in [
        "The repository does not yet include labelled assay or production grade data needed to train and validate Random Forest or XGBoost models.",
        "The reported model metrics in the current prototype are demonstration values from the estimator code, not final field-validated accuracy results.",
        "The Geosoft GRD binary parser should be independently validated against known grid values before scientific conclusions are drawn from raw values.",
        "Some outside-study-area geological and environmental values are simulated or rule-based, so they should be treated as exploratory context only.",
        "AI-generated geological interpretation should support, not replace, field mapping, sampling, laboratory assays, and professional resource estimation.",
    ]:
        add_bullet(doc, item)
    add_paragraph(
        doc,
        "These limitations do not reduce the value of the prototype; rather, they "
        "define the next stage of research. The present system proves that the "
        "workflow can be implemented, deployed, and demonstrated. The next stage is "
        "to improve the scientific reliability of the prediction engine through "
        "validated training data and rigorous model comparison."
    )

    add_heading(doc, "15. Conclusion and Recommendations", 1)
    add_paragraph(
        doc,
        "ARUM ALPHA demonstrates how mineral estimation for Bukuru and the Jos Plateau "
        "can be supported with a modern web dashboard, geospatial visualization, "
        "radiometric feature engineering, and AI-assisted interpretation. The project "
        "successfully converts the original proposal into a usable prototype that can "
        "guide exploratory analysis and communicate mineral potential to students, "
        "supervisors, and technical reviewers."
    )
    add_paragraph(
        doc,
        "The expanded report shows that the project has both an academic and practical "
        "component. Academically, it is grounded in mineral prospectivity, radiometric "
        "data interpretation, geostatistics, and machine learning. Practically, it "
        "delivers a deployed web application that reviewers can open and inspect. "
        "This combination makes the project stronger than a purely theoretical study "
        "because it produces a working artifact while also explaining the scientific "
        "basis and limitations of the method."
    )
    add_heading(doc, "Recommendations for Further Work", 2)
    for item in [
        "Acquire verified assay or production grade data for tin, then train Random Forest and XGBoost regression models.",
        "Compare supervised models against kriging, inverse distance weighting, and linear regression baselines using RMSE, MAE, R2, and MAPE.",
        "Validate the GRD parser with known Geosoft/Oasis Montaj exports and document the exact grid dimensions and null-value handling.",
        "Add geological layers such as faults, lithology, DEM, slope, aspect, and historical mine locations.",
        "Implement persistent storage for estimation history and field verification results.",
    ]:
        add_numbered(doc, item)


def add_comprehensive_final_chapters(doc: Document):
    add_heading(doc, "Expanded Final Project Chapters", 1)
    add_paragraph(
        doc,
        "This expanded section provides the fuller academic and technical discussion "
        "expected in a final project file. It extends the proposal-style material into "
        "a complete project narrative by explaining the geological basis, data logic, "
        "software design, validation approach, deployment value, limitations, and "
        "future research direction in greater detail."
    )

    sections = [
        (
            "Expanded Geological and Mining Context",
            [
                "The Bukuru and Jos South area is important to this project because it represents a real mineral district rather than an imaginary case study. Tin mining around the Jos Plateau has influenced settlement, infrastructure, local livelihoods, and the development of mining engineering knowledge in Nigeria. A final project on mineral estimation should therefore explain both the technical and local context. The dashboard is not only a software exercise; it is a response to the long-standing need for better exploration planning in an area where mining history, abandoned workings, and complex geology all affect decision-making (Sheet 168 Naraguta Radiometric Metadata, 2021).",
                "The Younger Granite province provides a useful geological framework for the project. Cassiterite mineralization is often linked with specialized granites, greisenization, hydrothermal alteration, and weathering processes that concentrate resistant minerals. These processes do not always produce simple spatial patterns. A point with favourable radiometric readings may be close to mineralized rock, but it may also reflect lithological background, weathered material, or transported sediment. The report therefore treats radiometric interpretation as an indicator of prospectivity rather than a direct measurement of ore reserve.",
                "One reason machine learning is attractive in this setting is that mineralization is controlled by several overlapping factors. Lithology, structure, alteration, topography, drainage, and human mining history can all influence where tin is found. Traditional mapping remains essential, but it can be strengthened by computational systems that organize evidence and reveal spatial patterns. ARUM ALPHA uses this idea by connecting coordinates, radiometric data, and AI-assisted interpretation in one interface. The project demonstrates how geological reasoning can be made more accessible without removing the need for field verification.",
                "The study area also has educational value. Because Bukuru and Naraguta are recognizable locations within Plateau State, the project can be explained to local reviewers in concrete terms. The map is not abstract; it shows an area with known mining relevance. This makes it easier to defend the project because the technical outputs can be related to familiar communities, roads, and geological history. A dashboard that opens to the actual study area helps bridge the gap between written theory and practical exploration planning.",
                "The final report must avoid overstating what the prototype can prove. A deployed application can show mineral potential patterns, but it cannot certify a reserve without sampling, laboratory assays, and professional resource classification. For that reason, the language of the system uses terms such as predicted potential, confidence, and risk. These terms are appropriate for a decision-support system. They encourage careful interpretation and help prevent the misunderstanding that a software prediction is the same as a measured ore grade.",
                "The geological context therefore shapes the entire project. It determines why K, Th, and U data were selected, why the map focuses on Sheet 168 Naraguta, why tin is the target commodity, and why future work must include field data. Without this context, the system would look like a generic map application. With the context, it becomes a localized mining engineering tool designed for a specific mineral province and a specific exploration problem.",
            ],
        ),
        (
            "Expanded Radiometric Data Interpretation",
            [
                "Radiometric data are valuable because potassium, thorium, and uranium occur naturally in many rocks and can be measured remotely through gamma-ray surveys. In granitic terrains, changes in these elements may indicate differences in mineral composition, alteration, weathering, or structural control. For the ARUM ALPHA project, the data layers provide a starting point for identifying areas that deserve closer inspection. The system does not claim that a high radiometric reading automatically means high tin grade. Instead, it uses radiometric behaviour as part of a broader interpretation workflow.",
                "Potassium can be associated with feldspar-rich rocks and alteration minerals. In granite-related mineral systems, elevated potassium may reflect lithological variation or alteration processes that affected the host rock. However, potassium alone is not a reliable tin predictor. It must be interpreted together with thorium, uranium, geological setting, and surrounding spatial patterns. The dashboard helps with this by presenting K values alongside other variables instead of treating one measurement as the entire basis for a decision.",
                "Thorium is important in the prototype because elevated thorium can point to radiometric anomalies in granitic environments. The current estimator uses a rule-based threshold where Th greater than 20 ppm contributes to higher potential. This threshold is not a final scientific model; it is a transparent prototype rule based on the assumption that anomalous thorium may indicate favourable geological conditions. The rule is useful for demonstration because it gives supervisors a clear example of how geoscience knowledge can be converted into software logic.",
                "Uranium is treated carefully because its distribution may be affected by mobility, alteration, and environmental conditions. The project uses uranium both as an individual value and as part of ratio-based indicators. The Th/U ratio is especially useful because it expresses the relationship between two radiometric elements rather than looking at each in isolation. A high ratio may suggest relative thorium enrichment, which can be relevant in differentiated granitic systems. Again, the ratio is a guide rather than a final proof of ore grade.",
                "The K/(Th+U) ratio provides another way to reduce the complexity of the raw data. Instead of presenting three separate values only, the system calculates an index that balances potassium against thorium and uranium. Such feature engineering is common in machine-learning workflows because models often perform better when raw measurements are transformed into variables that capture relationships. In the final supervised model, additional ratios and spatial features could be tested to determine which variables have the strongest predictive power.",
                "A major limitation of the current data stage is the need for validated raw grid parsing. The repository includes Geosoft GRD files and XML metadata, but binary grid formats can be difficult to parse correctly without specialist tools. The current application includes fallback logic to keep the dashboard functional for demonstration. In a final scientific version, the grid values should be checked against known exports from Oasis Montaj or another verified GIS/geophysical package. This validation step would strengthen the credibility of every prediction made by the system.",
            ],
        ),
        (
            "Expanded Machine Learning and Geostatistical Foundation",
            [
                "The project proposal identified Random Forest and XGBoost as future supervised learning models because both methods are strong for nonlinear tabular prediction. Mineralization rarely follows a single linear relationship. A location may be prospective because of combined effects from thorium, uranium, potassium, distance to structure, elevation, drainage, and historical workings. Tree-based ensemble models can handle such interactions better than simple linear regression. This makes them appropriate candidates once labelled grade data are available (Breiman, 2001; Chen and Guestrin, 2016).",
                "Random Forest is useful because it builds many decision trees and combines their outputs. Each tree can learn different patterns from the data, and the ensemble reduces the risk of relying on one unstable model. In a mineral estimation project, Random Forest can also provide feature importance values, which help explain whether thorium, potassium, ratios, or spatial variables have stronger influence. Such interpretability matters in a school project because the student must not only present predictions but also explain why the model behaves as it does (Breiman, 2001).",
                "XGBoost is useful because it builds trees sequentially, with each new tree correcting errors from previous trees. This can produce high accuracy when the dataset is well prepared. It can also handle missing values and complex relationships. For ARUM ALPHA, XGBoost would be appropriate after obtaining tin grade labels from assays or historical production records. The final model could compare XGBoost against Random Forest, linear regression, and geostatistical baselines to determine whether machine learning truly improves prediction for the study area (Chen and Guestrin, 2016).",
                "Geostatistics remains important because mineral data are spatial. Ordinary kriging, inverse distance weighting, and variogram analysis provide methods for interpolating values across unsampled areas. These approaches should not be dismissed simply because machine learning is modern. Instead, a strong final project should compare geostatistical and machine-learning approaches. If machine learning performs better, the report should show this using metrics. If geostatistics performs similarly, that result is also useful because it identifies the simplest reliable method (Goovaerts, 1997).",
                "The present prototype uses AI-assisted interpretation rather than a trained supervised model. This design decision was practical because the repository does not contain labelled tin assay data. The AI model receives structured geological context and returns an interpretation, while the fallback estimator applies transparent rules. This gives the project a working prediction interface, but the report clearly explains that final scientific validation must come later. This honesty is important because it protects the project from overclaiming.",
                "A complete machine-learning workflow would include data cleaning, exploratory data analysis, feature engineering, train-validation-test splitting, model training, hyperparameter tuning, model comparison, uncertainty assessment, and spatial output generation. ARUM ALPHA already implements several early parts of this workflow: data organization, feature calculation, map visualization, and prediction response formatting. The missing piece is a labelled target variable. Once that target is acquired, the existing architecture can be extended into a full supervised ML system.",
            ],
        ),
        (
            "Expanded Data Processing Pipeline",
            [
                "The data processing pipeline begins with the local grid files and their metadata. The XML files provide the spatial extent, coordinate system, projection, and data creation details. This metadata is essential because a grid value is meaningful only when its location is known. The application uses the WGS84 / UTM Zone 32N context to relate projected coordinates to latitude and longitude. Without correct coordinate handling, the map would show points in the wrong location, making the entire interpretation unreliable.",
                "After loading the grid information, the system combines potassium, thorium, and uranium values into a common radiometric record. Each record includes projected coordinates, geographic coordinates, and the three radiometric measurements. This combined structure is necessary because the estimator needs all variables at once. It also allows the map to display each point with a colour based on radiometric potential. The design is simple, but it reflects an important data integration step in geoscience workflows.",
                "The system includes nearest-point search because users may click anywhere on the map or type coordinates that do not exactly match a grid node. Nearest-neighbour lookup provides a practical way to connect user input with available data. The estimator can also use surrounding points within a radius, which gives a local context instead of depending on a single point. This is important because geophysical data are spatially continuous, and one isolated measurement may be less reliable than a neighbourhood pattern.",
                "Feature calculation occurs after spatial lookup. The system calculates K/(Th+U), Th/U, and summary averages for nearby points when available. These calculated features are passed into the AI prompt and fallback logic. This approach demonstrates that machine learning is not only about choosing an algorithm; it also depends on the quality and meaning of input variables. A model cannot learn useful patterns if the features do not represent geological relationships.",
                "The API response is designed to support both display and export. It includes the prediction object, surrounding predictions, model metrics, analysis text, and recommendations. This structure is important because it separates raw values from interpretation. The dashboard can display the grade, confidence, and risk visually, while the exported JSON can preserve the full result for later review. In future work, the same response structure could be saved to a database.",
                "A strong future pipeline would include data provenance tracking. Each value should indicate whether it came from a measured grid, synthetic fallback data, online geological context, or user input. Provenance matters because different sources have different reliability. If the dashboard clearly labels data source and confidence, users can make better decisions and avoid treating preliminary estimates as final measurements. This is especially important in a mining context where decisions can have financial and environmental consequences.",
            ],
        ),
        (
            "Expanded Software Architecture and Implementation",
            [
                "The software architecture uses Next.js because it supports both frontend pages and backend API routes in one project. This makes it suitable for a student project where development speed and deployment simplicity matter. The dashboard is built with React components, while API routes handle data loading and estimation. This organization keeps the codebase understandable: interface logic stays in components, data logic stays in library files, and request handling stays in the app/api directory (ARUM ALPHA Project Repository, 2026).",
                "The Dashboard component acts as the main controller for the user interface. It stores selected coordinates, prediction results, loading state, search input, map layer state, and export behaviour. This component connects user actions to the API. For example, when a user enters coordinates and clicks the estimation button, Dashboard sends a POST request to /api/estimate and then passes the returned prediction to the PredictionPanel. This structure shows clear separation between interaction, computation, and presentation.",
                "The MapView component handles the Leaflet map. It displays base map layers, the study area rectangle, radiometric data points, prediction markers, selected point markers, and the optional heatmap overlay. Dynamic import is used because Leaflet depends on browser-specific APIs and can cause server-side rendering issues. This is a practical implementation detail that demonstrates awareness of the difference between server and client execution in modern web frameworks.",
                "The PredictionPanel component is responsible for communicating results. It presents predicted grade, confidence, risk level, K/Th/U readings, ratios, location information, AI analysis, recommendations, and model metrics. This component is important because mineral estimation results must be understandable. A raw number is less useful than a structured explanation. The panel turns model output into a readable decision-support summary.",
                "The grdParser library contains geospatial and data-processing functions. It reads files, applies coordinate assumptions, converts UTM to latitude/longitude, combines radiometric layers, generates synthetic fallback data, and finds nearest or surrounding points. This file is the technical bridge between geophysical data and the web dashboard. Its future improvement is one of the most important recommendations because scientific reliability depends heavily on correct data parsing.",
                "The groq library contains the AI estimation logic. It builds a prompt that includes location, UTM coordinates, K/Th/U values, K/(Th+U), Th/U, and geological background. The expected AI response is JSON, which makes it easier for the application to parse. The fallback estimator applies explicit rules when the AI service is unavailable. This fallback is good engineering practice because it prevents the application from becoming useless during API failure.",
            ],
        ),
        (
            "Expanded User Interface and Dashboard Design",
            [
                "The dashboard interface was designed around the typical exploration question: what is the mineral potential at this location? To answer that question, the user needs a map, a way to select or search for a location, a way to run estimation, and a clear result panel. The interface places controls on the left and the map on the right for desktop use. This layout supports quick scanning and keeps the spatial context visible while the user interacts with coordinate inputs.",
                "The search panel helps users find familiar locations such as Bukuru, Naraguta, Rayfield, Gyel, and Jos. This matters because many users think in terms of community names rather than coordinates. The coordinate panel supports more technical users who know latitude and longitude. By supporting both workflows, the dashboard becomes more inclusive. It can be used by students, supervisors, mining engineers, and non-specialist reviewers during a project demonstration.",
                "The map layer controls allow the user to show or hide prediction heatmaps. Layer control is important because geospatial visualizations can become crowded. A user may want to inspect base map features, radiometric points, or prediction markers separately. The current project includes basic layer functionality, but future versions could add lithology, faults, DEM, drainage, historical mine sites, and sampling locations as independent layers.",
                "The legend explains the colour interpretation for high, medium, and low potential sites. This is important for communication because colours can be misunderstood without context. In the project, red, yellow, and green markers help reviewers quickly see possible areas of interest. However, the report should clarify that these colours represent model output and not confirmed ore grade. A good dashboard must be visually clear and scientifically cautious at the same time.",
                "The mobile screenshot proves that the interface is responsive. In field-related work, mobile access can be useful because users may review maps outside the office. The current mobile layout stacks panels vertically, which is appropriate for a narrow screen. Future versions could improve mobile fieldwork by adding GPS location, offline map caching, saved observations, and photo attachment for field verification notes.",
                "The visual design uses a dark dashboard theme because it gives strong contrast for controls and makes map panels stand out. The design is suitable for a technical application rather than a marketing website. It focuses on tools, data, and results. This is appropriate for a mining engineering project because the main goal is interpretation, not decoration. The screenshots included in the report provide evidence that the deployed application is a real and inspectable product.",
            ],
        ),
        (
            "Expanded Testing, Validation, and Uncertainty Discussion",
            [
                "Testing in this project has two levels: software testing and scientific validation. Software testing checks whether the application loads, builds, displays data, responds to user input, and opens in production. Scientific validation checks whether predictions agree with real mineral grade evidence. The project has stronger evidence for software testing than scientific validation because the available repository does not include labelled assay data. The report keeps this distinction clear.",
                "The successful production build shows that the TypeScript and Next.js application can compile for deployment. Build success is important because it catches many syntax, import, and type-related issues. It does not guarantee scientific accuracy, but it confirms that the implementation is technically stable enough to run. The deployed Vercel link provides another test because the system loads outside the development environment.",
                "The screenshot test confirms that the dashboard can render in a browser and that map-based data points appear on the deployed page. This is a practical validation step for a project defence because the supervisor can compare the screenshots with the live application. If the live app matches the report, it increases confidence that the documentation is based on a real implementation rather than an imagined design.",
                "API testing should include valid coordinates, invalid coordinates, inside-study-area coordinates, and outside-study-area coordinates. For valid inside-area coordinates, the API should return radiometric data and prediction output. For invalid inputs, it should return clear error messages. For outside-area coordinates, it should use broader geological context or fallback logic. These cases are important because a robust system must respond gracefully to different user behaviours.",
                "Uncertainty should be communicated more strongly in future versions. Every prediction should include not only a grade estimate and confidence score but also an explanation of data source, distance to nearest known point, number of surrounding points used, and whether values were measured or synthetic. Such uncertainty indicators would make the system more trustworthy. They would also help users decide whether a location needs more field sampling before any decision is made.",
                "Scientific validation requires field data. The ideal validation dataset would include coordinates, assayed tin grade, lithology, alteration type, sample type, and quality-control information. With such data, the project could train supervised models and test whether radiometric variables truly predict grade. It could also compare results against kriging and IDW. Until that validation is completed, ARUM ALPHA should be described as a prototype and decision-support tool.",
            ],
        ),
        (
            "Expanded Deployment, Maintenance, and Sustainability Discussion",
            [
                "Deployment to Vercel is an important achievement because it turns the project from a local codebase into a public demonstration. A local application can fail to impress reviewers if it requires complex setup. A deployed application can be opened directly from a browser. This is useful for a final-year project because supervisors, classmates, and external examiners can inspect the system without installing Node.js or downloading the repository (ARUM ALPHA Deployed Application, 2026).",
                "The deployed system also creates responsibility. If a public link is shared, the project should avoid exposing private API keys or sensitive data. The use of environment variables for the Groq API key is therefore appropriate. In production, secrets should never be committed to GitHub. Future work could also add rate limiting, logging, and clearer error handling so that the application remains stable under repeated use.",
                "Maintenance will be necessary if dependencies change. Next.js, React, Leaflet, and other packages are updated frequently. A final project may work at submission time but later fail if deployment dependencies change or environment variables are removed. For that reason, the report should document installation steps, build commands, environment variables, and deployment assumptions. This makes the project easier to maintain or reproduce in the future.",
                "Sustainability is relevant because mineral exploration can affect land, water, communities, and livelihoods. A digital target-ranking tool can support sustainability by reducing unnecessary field disturbance. Instead of sampling randomly, users can prioritize locations with stronger evidence. However, digital predictions can also create risk if they are misused to justify exploration without proper permits or community consultation. The report therefore recommends responsible use and field verification.",
                "The system could be extended into a participatory data platform. Field teams could upload sample results, photographs, and site notes. Communities or regulators could view maps showing areas of interest and environmental sensitivity. Such a system would require careful governance, but it could improve transparency. In its current form, ARUM ALPHA is a prototype; in future, it could become a broader geoscience decision-support platform.",
                "A sustainable long-term version should also support data export, model versioning, and audit trails. Users should know which model generated a prediction, what data were used, and when the prediction was produced. This is especially important in mining because decisions may be reviewed later. An audit trail would help distinguish preliminary academic predictions from validated professional outputs.",
            ],
        ),
        (
            "Expanded Socioeconomic and Environmental Relevance",
            [
                "Mining projects have socioeconomic importance because they can create employment, generate revenue, and support local development. However, they can also create environmental damage, unsafe abandoned pits, land-use conflict, and community concerns. A mineral estimation system should therefore not be seen only as a technical tool. It is part of a wider decision-making environment where economic opportunity must be balanced with environmental responsibility.",
                "In Plateau State, tin mining history has left both positive and negative legacies. The region is known for mineral wealth, but historic mining has also contributed to disturbed land and abandoned workings. A modern exploration system should learn from this history. By improving early-stage target selection, ARUM ALPHA can help reduce unnecessary disturbance. Better information can support better planning, but only if users interpret the outputs responsibly.",
                "The project may also support local capacity building. Many students in mining engineering learn about geostatistics, exploration, and ore grade estimation in theory, but they may not always build digital tools. ARUM ALPHA shows that students can connect mining knowledge with web development, data processing, and AI systems. This combination is valuable because the future mining industry will increasingly require digital competence.",
                "Environmental planning can be integrated into future versions of the dashboard. Layers such as drainage, protected areas, settlements, farmland, erosion risk, and abandoned mine hazards could be added. If mineral potential is high but environmental sensitivity is also high, the system could flag the location for careful review. This would make the dashboard more balanced and more useful for sustainable exploration planning.",
                "Community engagement is another future consideration. A technical prediction may be accurate, but exploration still affects people who live near the target area. Future project development could include a communication module explaining what the maps mean in plain language. This would reduce misunderstanding and support responsible engagement. The report therefore recommends that technical outputs should always be paired with ethical communication.",
                "The socioeconomic relevance of the project is strongest when the system is understood as a support tool rather than a final authority. It can help identify areas for further study, organize data, and communicate evidence. It cannot replace professional judgement, regulatory approval, or community consultation. This balanced interpretation makes the project more mature and more suitable for final submission.",
            ],
        ),
        (
            "Expanded Future Work and Research Roadmap",
            [
                "The most important future work is the acquisition of labelled grade data. Without target labels, the system cannot train a true supervised model. Suitable labels could come from assay results, historical production data, verified sample databases, or supervised field campaigns. Once labels are available, the project can move from AI-assisted interpretation to statistically evaluated ore grade prediction.",
                "The second priority is validation of the GRD parser. Geophysical file formats can be complex, and incorrect parsing would affect every downstream result. The project should compare parser outputs with values exported from trusted geophysical software. If differences are found, the parser should be corrected before model training. This step is technical but essential for scientific credibility.",
                "The third priority is integration of additional geospatial layers. Radiometric data alone may not capture all mineralization controls. Faults, lithology, slope, elevation, drainage, soil geochemistry, magnetic data, and historical mine locations could improve model performance. These layers would also make the dashboard more useful for exploration planning because users could compare multiple forms of evidence in one interface.",
                "The fourth priority is implementation of a database. At present, prediction results exist mainly in browser state and exported JSON. A database would allow users to save locations, compare multiple estimates, store field verification outcomes, and build a growing dataset for training. This would transform the prototype into a living exploration information system. A simple PostgreSQL or spatial database could support this future stage.",
                "The fifth priority is improved uncertainty visualization. Instead of only showing potential classes, the map could show confidence surfaces, distance-to-data indicators, or uncertainty bands around predictions. This would help users understand where the model is strong and where it is weak. In mineral exploration, uncertainty is not a weakness; it is information that guides better sampling.",
                "The final research roadmap is to compare models and publish results. A complete final version should train Random Forest, XGBoost, linear regression, kriging, and IDW baselines using the same dataset. It should report RMSE, MAE, R2, MAPE, and spatial validation plots. The best model should then be deployed in the dashboard with clear documentation. This would turn ARUM ALPHA from a prototype into a stronger applied research output.",
            ],
        ),
    ]

    for title, paragraphs in sections:
        add_heading(doc, title, 2)
        for paragraph in paragraphs:
            add_paragraph(doc, paragraph)


def add_final_report_expansion(doc: Document):
    add_heading(doc, "Detailed Final Project Expansion", 1)
    add_paragraph(
        doc,
        "This chapter expands the final report beyond the core implementation summary. "
        "It explains the project as a complete engineering work: why the problem is "
        "important, how the requirements were translated into software features, how "
        "the data were treated, how the system should be evaluated, and how the work "
        "can be defended as both a mining engineering project and a computing project. "
        "Table 19 provides a requirement traceability matrix linking requirements to "
        "implemented evidence."
    )

    traceability_rows = [
        ("R1", "Display local study area", "The dashboard opens on the Bukuru/Naraguta region and shows the Sheet 168 study boundary with sampled radiometric points."),
        ("R2", "Accept coordinate input", "Latitude and longitude fields allow a user to test a specific point rather than only viewing a static map."),
        ("R3", "Support location search", "Named locations such as Bukuru, Naraguta, Rayfield, Gyel, Du, and Zawan can be searched for quick demonstration."),
        ("R4", "Estimate mineral potential", "The /api/estimate route combines radiometric context, engineered ratios, AI interpretation, and fallback rules."),
        ("R5", "Explain each result", "The PredictionPanel reports grade, confidence, risk, radiometric values, model notes, analysis, and recommendations."),
        ("R6", "Deploy for assessment", "The Vercel production link makes the application accessible to supervisors without local installation."),
        ("R7", "Document limitations", "The report separates prototype predictions from certified reserves and recommends assay-based validation before professional use."),
        ("R8", "Allow future extension", "The code organization leaves separate places for map rendering, data parsing, AI logic, API routes, and interface components."),
    ]
    add_table(doc, ["ID", "Requirement", "Implemented Evidence"], traceability_rows, [900, 2300, 6160], font_size=8.7)

    sections = [
        (
            "Problem Definition in Practical Terms",
            [
                "The central problem addressed by ARUM ALPHA is the difficulty of turning scattered geological and radiometric information into a clear decision-support output. In many student projects, the data, the map, the model, and the explanation are presented separately. This makes it difficult for a reviewer to see how a coordinate becomes an interpreted mineral potential result. The project solves this by bringing the map, the input point, the local radiometric values, the estimation logic, and the written explanation into one web-based workflow (ARUM ALPHA Project Repository, 2026).",
                "In practical mining exploration, early decisions are often made under uncertainty. A team may know that an area has tin mining history, but it may not know exactly which locations deserve immediate sampling. The project does not claim to remove that uncertainty. Instead, it organizes early evidence so that field verification can be planned more intelligently. This is a realistic contribution because exploration commonly begins with screening and prioritization before expensive field campaigns are carried out.",
                "The problem is also educational. Mining engineering students are expected to understand geology, grade estimation, geostatistics, economics, and environmental responsibility, but modern practice increasingly requires digital tools. ARUM ALPHA demonstrates how a student can connect these areas. The final report therefore explains not only what the software displays, but why each display item is meaningful to an exploration workflow. This makes the work easier to defend during assessment.",
            ],
        ),
        (
            "Rationale for Selecting a Web Dashboard",
            [
                "A web dashboard was selected because it is accessible, visual, and easy to demonstrate. A spreadsheet could store radiometric values, but it would not communicate spatial patterns effectively. A desktop GIS could display layers, but it may require installation and technical setup. A deployed web dashboard offers a middle position: it is easier to open than desktop GIS, but still capable of showing an interactive map, coordinate input, estimation output, and explanatory results.",
                "The deployed link is important because it proves that the project has moved beyond a private development environment. A final-year project can appear incomplete if it only runs on the developer machine. Deployment shows that the application was built, hosted, and tested in a production-like environment. It also makes the project easier for supervisors to inspect because the assessment can include both the written report and the live system.",
                "The dashboard format also supports repeated exploration. A user can search one location, inspect the result, move to another location, compare confidence and risk, and export results. This repeated interaction is closer to real exploration screening than a one-time screenshot. The report therefore treats the dashboard as the main project product, while the written document explains the engineering, geological, and methodological reasoning behind it.",
            ],
        ),
        (
            "Detailed Requirement Analysis",
            [
                "The first functional requirement is map-based visualization. The system must show the study area and its radiometric sample points because mineral estimation is inherently spatial. A prediction without a map can be difficult to interpret. The map helps users see whether a selected point is inside the local Naraguta grid, near other points, or outside the study area. This spatial awareness improves the quality of interpretation.",
                "The second requirement is coordinate-based estimation. A mining engineer or student may receive a proposed sampling point as latitude and longitude. The system therefore needs to accept direct coordinate input instead of forcing the user to search by place name only. Direct coordinate input makes the project more technical and more aligned with field survey practice. It also makes testing easier because the same coordinate can be repeated during demonstration.",
                "The third requirement is explanatory output. The system should not simply display a predicted grade. It should also explain confidence, risk, radiometric values, ratios, and recommendations. This requirement is important because a black-box number is weak evidence in a final project. Supervisors need to understand what variables were considered and what the output means. The PredictionPanel meets this requirement by presenting numerical and narrative information together.",
            ],
        ),
        (
            "Non-Functional Requirement Analysis",
            [
                "Usability is a major non-functional requirement. A project dashboard used during defence must be understandable within a short time. The interface therefore places the search and coordinate controls near the map and makes the result panel readable. It avoids requiring the user to run commands before seeing the main output. This usability decision supports both technical review and oral presentation because the student can demonstrate the system quickly.",
                "Reliability is another requirement. The application includes fallback estimation logic so that the main workflow can continue even if the AI service is unavailable. This matters because a live demonstration can fail if it depends completely on an external API. The fallback does not replace a fully trained model, but it provides continuity and makes the system more resilient. This is a reasonable engineering decision for a student project with limited infrastructure.",
                "Maintainability is also important. The project separates map rendering, dashboard state, data parsing, online geological context, AI estimation, and API routing into different files. This organization makes it easier to improve one area without breaking the whole application. For example, a future student could improve the GRD parser without rewriting the map component. Maintainability gives the project a stronger future beyond the current report.",
            ],
        ),
        (
            "Data Quality and Preprocessing Detail",
            [
                "The quality of mineral estimation depends strongly on the quality of the input data. The project uses potassium, thorium, and uranium radiometric grids linked to Sheet 168 Naraguta. These layers are useful because they provide continuous geophysical evidence over the study area. However, the report must recognize that radiometric evidence is indirect. It can indicate geological variation and possible alteration, but it cannot prove the existence of economic tin ore without ground truth (Sheet 168 Naraguta Radiometric Metadata, 2021).",
                "Preprocessing begins with understanding coordinate reference information. The dataset metadata indicates WGS 84 / UTM Zone 32N, which means raw coordinates must be handled carefully before being displayed on a latitude/longitude web map. A mistake in projection handling would place points in the wrong area and invalidate the interpretation. The report therefore emphasizes coordinate conversion as a core technical task rather than a minor programming detail.",
                "A complete preprocessing pipeline would also remove no-data values, check grid alignment, compare layer extents, calculate descriptive statistics, inspect outliers, and document data provenance. The current prototype performs enough transformation to support the dashboard demonstration, but future scientific validation should include a formal data audit. This is especially important if the application is later used to support field sampling or academic publication.",
            ],
        ),
        (
            "Feature Engineering Explanation",
            [
                "Feature engineering converts raw variables into more informative indicators. In ARUM ALPHA, potassium, thorium, and uranium are not only displayed as separate values; they are also used to calculate ratios such as K/(Th+U) and Th/U. These ratios are useful because mineral systems often depend on relationships among elements rather than isolated values. A ratio can highlight enrichment or depletion patterns that may not be obvious from raw readings alone.",
                "The Th/U ratio is especially useful in granitic terrains because it expresses relative thorium enrichment compared with uranium. Since uranium can be more mobile under some geological conditions, the ratio may provide clues about alteration or geochemical behaviour. In the prototype, this ratio is used as one of the explanatory features passed into the estimation prompt and fallback logic. The report describes it carefully as an indicator, not as a confirmed ore-grade measurement.",
                "Feature engineering also includes spatial context. A selected coordinate is more meaningful when nearby data points are considered. The system can find surrounding points and use their average behaviour to support interpretation. This reflects a basic principle of geoscience: one point should not be interpreted in isolation when neighbouring measurements are available. Future versions can improve this by adding distance-weighted features, local variance, and neighborhood anomaly scores.",
            ],
        ),
        (
            "Model Logic and Interpretability",
            [
                "The current system uses AI-assisted interpretation and a transparent fallback estimator. This approach was chosen because labelled tin assay data are not yet available in the repository. A supervised model requires known input-output pairs, such as radiometric features matched with measured tin grade. Without those labels, it would be misleading to claim that the system has trained a final predictive model. The report therefore describes the current model honestly as a prototype estimation workflow.",
                "Interpretability is important because mineral estimation can affect expensive decisions. A user needs to know why a location is rated high, medium, or low. The fallback logic provides this interpretability by using explicit threshold rules. The AI layer adds richer explanatory language, but the rule-based layer makes the assumptions visible. This combination supports project defence because the student can explain both the software behaviour and the geological reasoning behind it.",
                "In the next research stage, Random Forest and XGBoost can be trained once labelled grade data are available. These methods should be compared with geostatistical baselines instead of being assumed superior. A strong final model would report accuracy metrics, feature importance, cross-validation results, and spatial validation maps. Until then, ARUM ALPHA is best understood as a working decision-support prototype with a clear path toward supervised learning (Breiman, 2001; Chen and Guestrin, 2016; Goovaerts, 1997).",
            ],
        ),
        (
            "Interface Workflow Evaluation",
            [
                "The interface workflow begins when a user opens the deployed dashboard. The user can inspect the visible map, search for a location, enter coordinates, or click on the map. This design supports different user habits. Some reviewers may want to see known place names, while technical users may prefer exact coordinates. Providing both options makes the dashboard more flexible and helps the project demonstration flow smoothly.",
                "After location selection, the user runs the estimation workflow and receives a structured response. This response includes predicted tin grade, confidence, risk, radiometric readings, derived ratios, analysis, recommendations, and model metrics. The layout is useful because it separates the result into categories. A user can quickly read the headline estimate, then inspect the supporting evidence. This is better than showing a long unstructured paragraph without numeric context.",
                "The interface also supports visual comparison. Markers, colours, legends, and map layers help users compare areas of higher and lower predicted potential. This is important because mineral exploration is rarely about one point only. It is usually about comparing several possible targets and deciding which ones deserve further sampling. The dashboard makes that comparison easier than a static report table.",
            ],
        ),
        (
            "Testing Strategy in More Detail",
            [
                "Testing should begin with basic application behaviour. The development server should load without console errors, the production build should complete, and the deployed Vercel page should open in a browser. These tests prove that the software is runnable. They do not prove geological accuracy, but they are necessary because an accurate idea cannot be assessed if the system itself fails to run.",
                "API testing should include normal inputs, boundary inputs, and invalid inputs. Normal inputs include coordinates inside Bukuru or Naraguta. Boundary inputs include points close to the edge of the Sheet 168 study area. Invalid inputs include missing coordinates, text instead of numbers, and coordinates far outside the intended area. A robust system should return clear messages instead of crashing or producing confusing output.",
                "Scientific testing requires a different approach. Predictions should be compared with known assay values, historic workings, or validated sample points. The model should be evaluated with metrics such as mean absolute error, root mean squared error, R2, and classification accuracy if potential classes are used. Until this scientific testing is completed, the current prototype should be defended as a functioning estimation and visualization system rather than a certified reserve model.",
            ],
        ),
        (
            "Result Interpretation and Discussion",
            [
                "The results of this project should be interpreted at three levels. At the software level, the result is a deployed web application that opens, displays a map, accepts inputs, and returns structured estimates. At the data level, the result is an organized use of K, Th, and U radiometric information for local study-area interpretation. At the academic level, the result is a demonstration of how mining engineering knowledge can be integrated with modern web development and AI services.",
                "The most visible result is the dashboard screenshot included in the report. It shows the deployed interface, study-area map, location controls, and radiometric points. This evidence is important because it proves that the project produced a user-facing artifact. It also allows the written report to connect directly with the live implementation. A reader can move from the report to the production link and inspect the same system.",
                "The results should not be overstated. The application estimates potential and provides decision-support information; it does not classify mineral reserves under professional reporting codes. The absence of labelled tin assay data limits scientific validation. This limitation is not a failure if it is properly acknowledged. A mature final report explains what the project achieved, what remains uncertain, and what evidence would be needed to strengthen it.",
            ],
        ),
        (
            "Risk Analysis",
            [
                "The first project risk is data reliability. If the GRD values are parsed incorrectly or if fallback synthetic values are mistaken for measured data, the prediction could mislead users. The mitigation is to document data provenance clearly and validate parsed values against a trusted geophysical software export. The report therefore recommends raw grid verification as a priority future task.",
                "The second risk is model overconfidence. AI-generated explanations can sound convincing even when input data are limited. To reduce this risk, the dashboard includes confidence and risk fields and the report uses careful language. A future version should go further by showing model uncertainty, nearest-data distance, number of supporting points, and whether a result depends on measured or synthetic values.",
                "The third risk is deployment dependency. The live application depends on hosting configuration, package versions, and environment variables. If any of these change, the deployed link could fail. The mitigation is to document installation steps, maintain dependency versions, and keep the source code organized. The final report includes deployment notes so that the project can be reproduced or restored if the hosting environment changes.",
            ],
        ),
        (
            "Ethical and Environmental Considerations",
            [
                "Mineral prediction systems should be used responsibly. A map that highlights a potential target may influence land-use decisions, field visits, and community expectations. For this reason, the report emphasizes that ARUM ALPHA is a decision-support prototype. It should guide further investigation, not replace professional judgement, regulatory approval, or community consultation. This ethical boundary is important in a mining engineering project.",
                "Environmental responsibility is also relevant. The Jos Plateau has experienced historic mining disturbance, and future exploration should avoid repeating harmful practices. A digital screening tool can reduce unnecessary disturbance by helping users prioritize areas before fieldwork. However, better target selection does not remove the need for environmental assessment. Future versions should include layers for drainage, settlements, land use, erosion risk, and abandoned mine hazards.",
                "Community communication should be part of future development. Technical maps can be misunderstood by non-specialists, especially when they use strong colours such as red or green. A responsible system should explain that prediction classes are preliminary and require verification. The report recommends simple explanatory text, clear disclaimers, and responsible demonstration practices whenever the tool is shown outside a technical setting.",
            ],
        ),
        (
            "Academic Contribution of the Project",
            [
                "The academic contribution of ARUM ALPHA is the integration of mining engineering concepts with a deployed software product. Many projects discuss mineral estimation theoretically, while many software demos ignore the geological meaning of the data. This project connects both sides. It uses radiometric variables, geospatial coordinates, ratios, map layers, and AI interpretation to create a working tool that can be explained in mining engineering terms.",
                "The project also contributes a reusable structure for future student work. The same architecture can be adapted to other minerals, other study areas, and other geophysical datasets. If labelled data are later added, the estimation engine can be upgraded to train and compare supervised models. If environmental layers are added, the dashboard can become a broader exploration planning tool. This extensibility gives the work value beyond the immediate submission.",
                "Another contribution is communication. The dashboard turns technical values into a visual and narrative result that a supervisor can inspect. This is important because final-year projects are assessed not only on theory, but also on clarity of execution. ARUM ALPHA shows an end-to-end pathway from data and project idea to deployed artifact, screenshots, report documentation, and future research plan.",
            ],
        ),
        (
            "Limitations Expanded",
            [
                "The main limitation is the absence of verified tin assay labels in the project repository. Without labels, a final supervised ore-grade model cannot be trained or objectively validated. The current estimator can rank potential based on radiometric indicators and AI-assisted reasoning, but it cannot prove grade accuracy. This limitation should be stated clearly during project defence because it shows academic honesty and technical maturity.",
                "Another limitation is the reliance on radiometric variables alone. Tin mineralization can be influenced by lithology, structure, weathering, drainage, alteration, and historical mining activity. K, Th, and U provide useful clues, but a stronger model would include geological maps, fault distance, elevation, magnetic data, soil geochemistry, and known mine locations. The current project provides the software foundation for this integration, but the extra datasets are future work.",
                "A third limitation is that the system is a prototype rather than a regulated mining decision platform. It does not include user accounts, audit trails, formal database storage, professional resource classification, or field verification management. These features are not required for a school prototype, but they would be necessary for operational use. The report distinguishes the academic prototype from a full commercial system.",
            ],
        ),
        (
            "Implementation Reflection",
            [
                "The implementation process shows the value of incremental development. The project begins with data loading and coordinate handling, then adds map rendering, interface controls, estimation API routes, AI interpretation, fallback logic, deployment, screenshots, and documentation. This sequence is sensible because it builds from the data foundation toward the user-facing product. A dashboard without reliable data handling would be visually attractive but technically weak.",
                "The use of Next.js simplified the project because frontend pages and backend routes could live in one codebase. This reduces setup complexity for a student project. The dashboard components handle the interface while API routes handle estimation. This separation is not only cleaner for code maintenance; it also makes the report easier to explain because each file has a clear role in the final system.",
                "The final implementation is strongest as a demonstrable prototype. It gives a reviewer something to open, inspect, and test. It also creates a base that can be improved after submission. This is a reasonable outcome for a final-year project because the work demonstrates initiative, technical competence, and domain awareness while leaving a clear research path for deeper validation.",
            ],
        ),
        (
            "Applied Case Study Workflow",
            [
                "A practical case study workflow begins with the selection of a location in the Bukuru or Naraguta area. The user opens the deployed dashboard, confirms the map position, and either searches for a named location or enters coordinates obtained from a proposed field point. This first step is important because it connects the abstract report to a real spatial decision. The system is not only reading a dataset; it is helping a user ask a location-specific exploration question.",
                "After the point is selected, the dashboard retrieves the nearest radiometric context and prepares derived indicators. The user can then run the estimation workflow and inspect the predicted grade class, confidence, risk level, K, Th, U, and ratio values. In a field planning scenario, this output would not immediately authorize mining. Instead, it would help rank the point against other possible sampling locations so that limited field time can be used more efficiently.",
                "The final case study step is interpretation. The user should read the recommendation text, compare the point with surrounding data, and decide whether the location deserves field verification. A high-potential point with low confidence may require more sampling before any conclusion is made. A medium-potential point near known workings may still be interesting. A low-potential point may be useful as a control sample. This shows that the dashboard supports reasoning rather than replacing it.",
            ],
        ),
        (
            "Proposed Field Data Collection Plan",
            [
                "For the project to become a stronger scientific model, a structured field data collection plan is required. The first stage should involve selecting representative sampling locations across high, medium, and low predicted potential zones. Sampling only high-potential locations would create biased training data because the model would not learn what low-potential areas look like. A balanced sampling plan would give future supervised models a better chance of learning meaningful relationships.",
                "Each field sample should include more than a laboratory tin grade. The data sheet should record coordinates, elevation, sample type, lithology, alteration, weathering level, nearby structures, land-use condition, and field observations. Photographs can also be attached to each point. These supporting variables would help explain why a prediction succeeds or fails. They would also allow the dashboard to expand beyond radiometric interpretation into a broader geological evidence system.",
                "Quality control should be planned from the beginning. Duplicate samples, blank samples, reference standards, and coordinate checks can improve confidence in the dataset. If a future model is trained on poor field data, the model output will also be poor. The report therefore recommends that future field campaigns should follow a simple but disciplined quality-control procedure. This would make ARUM ALPHA more defensible as a research tool and not only as a demonstration dashboard.",
            ],
        ),
        (
            "Proposed Model Evaluation Plan",
            [
                "Once field or historical assay labels are available, the model evaluation plan should begin with exploratory data analysis. The student should plot grade distributions, inspect outliers, compare K, Th, U, and ratio values against grade, and check whether high-grade samples cluster spatially. This step prevents blind model training. It helps reveal whether the available data contain enough signal for prediction or whether additional variables are needed.",
                "The next step should compare several modelling methods under the same validation design. Random Forest and XGBoost should be compared with linear regression, inverse distance weighting, and ordinary kriging where appropriate. The comparison should use metrics such as mean absolute error, root mean squared error, coefficient of determination, and classification accuracy if grade classes are used. A model should be selected because it performs well and remains explainable, not because it sounds advanced (Breiman, 2001; Chen and Guestrin, 2016; Goovaerts, 1997).",
                "Spatial validation should also be included because normal random train-test splitting can exaggerate performance when nearby samples are similar. A stronger evaluation would hold out spatial blocks or test on locations that are separated from training points. This would show whether the model can generalize to new areas rather than only memorizing local patterns. For mineral exploration, this kind of validation is especially important because the goal is to predict unsampled locations.",
            ],
        ),
        (
            "Recommended System Enhancements",
            [
                "The first recommended enhancement is a database for saved predictions and field verification results. At present, the user can export JSON, but the system does not maintain a long-term project history. A database would allow users to save points, revisit previous estimates, attach assay results, and compare predicted potential with confirmed field evidence. This would turn the dashboard into a growing project archive rather than a single-session demonstration.",
                "The second enhancement is user authentication and role control. A future version could allow students, supervisors, and field officers to use different permissions. Students could save observations, supervisors could review verified results, and administrators could manage datasets. This is not necessary for the current school submission, but it would be useful if the system becomes a departmental research tool or a larger exploration planning platform.",
                "The third enhancement is richer map layering. Geological boundaries, fault lines, drainage, digital elevation models, abandoned mine sites, settlements, roads, and environmental sensitivity layers would make the dashboard more useful. These layers would help users interpret why a radiometric anomaly occurs and whether it is practical or responsible to investigate. A strong exploration decision should combine mineral potential with access, safety, environmental, and community information.",
            ],
        ),
        (
            "Defence and Demonstration Plan",
            [
                "During project defence, the student should begin by explaining the problem: early-stage tin potential assessment in the Bukuru/Naraguta area requires better integration of radiometric evidence and map-based interpretation. The student should then open the deployed production link and show the dashboard. This immediately proves that the project is implemented, deployed, and accessible. The report screenshots should match what is visible on the live page.",
                "The next demonstration step should be coordinate testing. The student can select a location inside the study area, run an estimate, and explain the output fields. The explanation should cover predicted grade, confidence, risk, K, Th, U, ratios, and recommendations. The student should also explain that the current model is a prototype and that final scientific validation requires labelled assay data. This balanced explanation will sound more credible than overclaiming.",
                "The final defence step should be future work. The student should describe how the project can be improved with field samples, lab assays, validated GRD parsing, geological layers, supervised learning, geostatistical comparison, uncertainty maps, and database storage. This shows that the student understands the difference between a working prototype and a professional estimation system. It also demonstrates that the project has a clear academic growth path.",
            ],
        ),
    ]

    for title, paragraphs in sections:
        add_heading(doc, title, 2)
        for paragraph in paragraphs:
            add_paragraph(doc, paragraph)


def add_references_appendix(doc: Document):
    add_heading(doc, "References", 1)
    refs = [
        "ARUM ALPHA Proposal Slides. (2026). Development of a machine learning based system for ore grade and mineral distribution prediction in Du, Bukuru, Jos South LGA, Plateau State, Nigeria. Local file: ARUM ALPHA PROPOSAL SLIDES.pptx.",
        "ARUM ALPHA Project Proposal. (2026). Development of a machine learning based system for ore grade and mineral distribution prediction in Du/Bukuru, Jos South LGA, Plateau State, Nigeria. Local file: ARUM_ALPHA_Project_Proposal.docx.",
        "ARUM ALPHA Project Repository. (2026). ARUM ALPHA source code, README, dashboard components, API routes, and local implementation files. Local repository: arum-alpha/.",
        f"ARUM ALPHA Deployed Application. (2026). ARUM ALPHA AI-powered mineral estimation dashboard. {DEPLOYED_URL}",
        "Sheet 168 Naraguta Radiometric Metadata. (2021). Sheet168_Naraguta_Potassium.grd.xml, Sheet168_Naraguta_Th.grd.xml, and Sheet168_Naraguta_U.grd.xml metadata files, dated April 16, 2021.",
        "Breiman, L. (2001). Random Forests. Machine Learning, 45, 5-32.",
        "Chen, T. and Guestrin, C. (2016). XGBoost: A scalable tree boosting system. Proceedings of KDD.",
        "Goovaerts, P. (1997). Geostatistics for Natural Resources Evaluation. Oxford University Press.",
        "Rodriguez-Galiano, V. F., Sanchez-Castillo, M., Chica-Olmo, M., and Chica-Rivas, M. (2015). Machine learning predictive models for mineral prospectivity. Ore Geology Reviews.",
    ]
    for ref in refs:
        add_bullet(doc, ref)

    add_heading(doc, "Appendix A: Running the Project", 1)
    add_paragraph(
        doc,
        "Prerequisites: Node.js 18 or newer and a Groq API key for live AI estimation. "
        "Table 20 lists the commands and values needed to run or inspect the project.",
    )
    commands = [
        ("Install dependencies", "npm install"),
        ("Configure environment", "copy env.example .env.local"),
        ("Start development server", "npm run dev"),
        ("Open application", "http://localhost:3000"),
        ("Open deployed production site", DEPLOYED_URL),
        ("Build for production", "npm run build"),
    ]
    add_table(doc, ["Task", "Command or Value"], commands, [3000, 6360])

    add_heading(doc, "Appendix B: Project Management Snapshot", 1)
    add_paragraph(
        doc,
        "Table 21 provides the estimated budget items used for the project management "
        "snapshot.",
    )
    budget_rows = [
        ("Data acquisition", "50,000"),
        ("Software and tools", "150,000"),
        ("Research assistance", "100,000"),
        ("Miscellaneous", "100,000"),
        ("Total", "400,000"),
    ]
    add_table(doc, ["Budget Item", "Estimated Cost (NGN)"], budget_rows, [4200, 5160])
    add_paragraph(
        doc,
        "The proposal timeline covers topic approval, literature review, data collection, "
        "model development, dashboard implementation, validation, final report writing, "
        "and presentation preparation between January and June 2026. Table 22 expands "
        "the timeline into monthly activities."
    )

    add_heading(doc, "Appendix C: Suggested Full Project Timeline", 1)
    timeline_rows = [
        ("January", "Topic selection, supervisor approval, preliminary background reading."),
        ("February", "Literature review, proposal refinement, and study-area definition."),
        ("March", "Data acquisition, metadata inspection, and project environment setup."),
        ("April", "Dashboard development, GRD parsing, map integration, and API route implementation."),
        ("May", "Prediction logic, AI integration, deployment testing, and report drafting."),
        ("June", "Final documentation, screenshots, presentation preparation, and project defence."),
    ]
    add_table(doc, ["Period", "Main Activity"], timeline_rows, [1800, 7560])

    add_heading(doc, "Appendix D: Future Data Collection Checklist", 1)
    for item in [
        "Tin assay results from verified samples across the study area.",
        "Historical production records from known Bukuru and Naraguta mining sites.",
        "Lithological boundary data and fault/lineament maps.",
        "Digital elevation model, slope, aspect, drainage, and accessibility layers.",
        "Ground radiometric readings for comparison with airborne or gridded data.",
        "Coordinates of abandoned pits, active exploration sites, and non-mineralized control locations.",
        "Environmental sensitivity and land-use information for responsible exploration planning.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Appendix E: Defence Demonstration Guide", 1)
    for item in [
        f"Open the deployed application at {DEPLOYED_URL}.",
        "Point out the study area label, data point count, coordinate input form, and map layers.",
        "Explain that the map shows sampled radiometric points from Sheet 168 Naraguta.",
        "Enter a coordinate within the study area and run the estimation workflow.",
        "Discuss the predicted grade, confidence, risk level, radiometric readings, and recommendations.",
        "Explain the difference between the current prototype and a future supervised Random Forest/XGBoost model.",
        "Close by identifying the validation data required for a publishable mineral prospectivity model.",
    ]:
        add_numbered(doc, item)


def audit_docx(docx_path: Path):
    import zipfile
    from xml.etree import ElementTree as ET

    with zipfile.ZipFile(docx_path) as z:
        doc_xml = ET.fromstring(z.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        tables = doc_xml.findall(".//w:tbl", ns)
        missing = []
        for idx, tbl in enumerate(tables, 1):
            tbl_w = tbl.find(".//w:tblPr/w:tblW", ns)
            grid_cols = tbl.findall(".//w:tblGrid/w:gridCol", ns)
            if tbl_w is None or tbl_w.get(qn("w:w")) is None or not grid_cols:
                missing.append(str(idx))
        if missing:
            raise RuntimeError(f"Tables missing fixed geometry: {', '.join(missing)}")


def build():
    global TABLE_COUNTER
    TABLE_COUNTER = 0
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)
    add_preliminary_sections(doc)
    add_abstract(doc)
    add_intro_problem_objectives(doc)
    add_lit_study_data(doc)
    add_design_method_impl(doc)
    add_deployment_and_screenshots(doc)
    add_results_limits_conclusion(doc)
    add_comprehensive_final_chapters(doc)
    add_final_report_expansion(doc)
    add_references_appendix(doc)
    doc.save(OUT)
    audit_docx(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
